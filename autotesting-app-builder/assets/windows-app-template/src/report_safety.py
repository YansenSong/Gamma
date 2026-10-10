"""Visible draft protection for generated copies; contains no scoring business rules."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

DRAFT_TITLE = '工程核对稿 — 非正式检验报告'

def requires_draft(spec):
    return (spec.get('release_status') != 'business_validated'
            or spec.get('validation_state', {}).get('status') != 'approved')

def mark_engineering_copy(path, policy):
    """Add visible first-page identity and persistent per-page provenance notices."""
    document = Document(path)
    title = policy.get('title', DRAFT_TITLE)
    provenance = policy['provenance']
    paragraph = document.paragraphs[0].insert_paragraph_before(title)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    for run in paragraph.runs:
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor.from_string('A61B1B')
    changed = []
    title_replacements = []
    body_replacements = []
    seen = set()
    for section in document.sections:
        minimum=policy.get('minimum_header_footer_distance_inches')
        if minimum is not None:
            section.header_distance=max(section.header_distance or 0,Inches(minimum))
            section.footer_distance=max(section.footer_distance or 0,Inches(minimum))
        for part in (section.header, section.first_page_header, section.even_page_header,
                     section.footer, section.first_page_footer, section.even_page_footer):
            if part.part.partname in seen:
                continue
            seen.add(part.part.partname)
            is_header = 'header' in str(part.part.partname)
            replaced=False
            if is_header:
                for table in part.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                replacement=policy.get('header_title_replacements',{}).get(paragraph.text.strip())
                                if replacement:
                                    title_replacements.append({'part':str(part.part.partname),'before':paragraph.text,'after':replacement})
                                    paragraph.text=replacement
                                    for run in paragraph.runs:
                                        run.bold=True;run.font.size=Pt(14);run.font.color.rgb=RGBColor.from_string('A61B1B')
                                    replaced=True
                if replaced:
                    changed.append(str(part.part.partname))
                    continue
            # Linked parts are edited once; original header tables remain intact.
            empty=next((p for p in part.paragraphs if not p.text.strip()),None)
            p = empty if empty is not None else part.add_paragraph()
            p.add_run(title if is_header else title+'\n'+provenance)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                run.bold = is_header
                run.font.size = Pt(8 if is_header else 7)
                run.font.color.rgb = RGBColor.from_string('A61B1B')
            changed.append(str(part.part.partname))
    # Empty signature boxes must also state that no approval occurred.
    def tables(items):
        for table in items:
            yield table
            for row in table.rows:
                for cell in row.cells:
                    yield from tables(cell.tables)
    signatures = []
    seen_cells = set()
    for table in tables(document.tables):
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                for p in cell.paragraphs:
                    if p.text.strip() in policy.get('unsigned_labels', []):
                        run=p.add_run('未审批')
                        run.font.size=Pt(8)
                        signatures.append(p.text)
    paragraphs=list(document.paragraphs)
    for table in tables(document.tables):
        for row in table.rows:
            for cell in row.cells:paragraphs.extend(cell.paragraphs)
    replacements=policy.get('body_text_replacements',{})
    for p in paragraphs:
        if p.text.strip() in replacements:
            before=p.text; p.text=replacements[before.strip()]
            body_replacements.append({'before':before,'after':p.text})
    document.save(path)
    return {'title': title, 'provenance': provenance, 'header_footer_parts': changed,'header_title_replacements':title_replacements,
            'unsigned_boxes': signatures, 'body_text_replacements':body_replacements,
            'original_page_count_text': 'preserved; not a verified page count'}
