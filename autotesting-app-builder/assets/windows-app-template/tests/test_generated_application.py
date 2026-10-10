"""Contract tests for this generated application.

This file ships as a skeleton. It states the checks the Skill requires of every
generated application, and marks what a real project has to supply. Until
`config/project.json` is filled in, every test here **skips**: a placeholder
project has no sources, no output cells and no report fields, so there is nothing
to score and nothing to check.

Fill in per project:
  - `config/project.json` — sources, scoring.inputs, excel_output, report_output
  - `templates/`           — the score workbook and the DOCX report template
  - `examples/input.json`  — one realistic input record
  - `examples/boundary_cases.json` (optional) — see the last test

Run, matching the command the upstream AutoTesting generated apps record:

    uv run --with pytest --with openpyxl --with docxtpl --with python-docx \\
        python -m pytest tests -q

The behavioural tests drive native Excel recalculation, so they need Microsoft
Excel on the host — the same requirement the application itself has.
"""
from __future__ import annotations
import hashlib, json, sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / "src"))

SPEC = json.loads((PROJECT / "config" / "project.json").read_text(encoding="utf-8"))
SOURCE_ROLES = {"charter", "workbook", "report_template"}


def _configured(spec: dict) -> bool:
    return (bool(spec.get("sources"))
            and bool(spec.get("report_output", {}).get("fields"))
            and bool(spec.get("excel_output", {}).get("output_cells")))


# A placeholder project cannot be scored. Skipping loudly beats failing on a
# configuration that is deliberately not filled in yet.
pytestmark = pytest.mark.skipif(
    not _configured(SPEC),
    reason="placeholder project: fill in config/project.json (sources, "
           "excel_output.output_cells, report_output.fields) before these tests "
           "can check anything")


def sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def service():
    """The application's own service module, imported lazily so the skip path
    stays clean on a project whose dependencies are not installed yet."""
    import app_service
    return app_service


@pytest.fixture(scope="module")
def example_input() -> Path:
    path = PROJECT / "examples" / "input.json"
    if not path.is_file():
        pytest.skip("no examples/input.json: add one realistic input record")
    return path


@pytest.fixture(scope="module")
def job(service, example_input, tmp_path_factory):
    """One scored run, shared by the behavioural tests below. Each run drives
    native Excel through PowerShell, so it is deliberately not repeated."""
    return service.run_job(example_input, tmp_path_factory.mktemp("run"))


# --- structure: true of any configuration, no Excel needed -------------------

def test_declares_the_three_required_sources():
    roles = {source["role"] for source in SPEC["sources"]}
    assert SOURCE_ROLES <= roles, f"missing source roles: {sorted(SOURCE_ROLES - roles)}"
    for source in SPEC["sources"]:
        path = PROJECT / source["path"]
        assert path.is_file(), f"source file is missing: {source['path']}"
        assert sha256(path) == source["sha256"], f"source hash drifted: {source['path']}"


def test_every_score_cell_holds_a_formula():
    # Excel formulas are the only scoring authority. A mapped score cell holding
    # a literal means someone typed a result instead of deriving it from the
    # charter, and the audit trail no longer proves anything.
    book = load_workbook(PROJECT / SPEC["excel_output"]["template"], data_only=False)
    try:
        for name, mapping in SPEC["excel_output"]["output_cells"].items():
            cell = book[mapping["sheet"]][mapping["cell"]]
            assert cell.data_type == "f", (
                f"{name} ({mapping['sheet']}!{mapping['cell']}) is not a formula")
    finally:
        book.close()


def test_report_template_markers_equal_declared_fields():
    from project_validation import docx_markers
    markers = docx_markers(PROJECT / SPEC["report_output"]["template"])["names"]
    declared = set(SPEC["report_output"]["fields"])
    assert not markers - declared, f"template markers no field declares: {sorted(markers - declared)}"
    assert not declared - markers, f"declared fields with no marker: {sorted(declared - markers)}"


def test_example_input_uses_only_declared_inputs(example_input):
    facts = json.loads(example_input.read_text(encoding="utf-8"))
    inputs = SPEC["scoring"]["inputs"]
    declared = {item["id"] for item in inputs}
    assert set(facts) <= declared, f"unknown inputs: {sorted(set(facts) - declared)}"
    required = {item["id"] for item in inputs if item.get("required")}
    assert not required - set(facts), (
        f"required inputs absent from the example: {sorted(required - set(facts))}")


# --- behaviour: needs a native Excel run ------------------------------------

def test_scoring_produces_a_score_for_every_output_cell(job):
    for output_id in SPEC["excel_output"]["output_cells"]:
        assert output_id in job["scores"], f"no score produced for {output_id}"


def test_rendered_report_has_no_unrendered_markers(job):
    from docx import Document
    document = Document(job["report"])
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "{{" not in text, "the report rendered with an unfilled {{ }} marker"
    assert "}}" not in text, "the report rendered with an unfilled {{ }} marker"


def test_audit_records_the_calculation_authority_and_output_hashes(job):
    audit = json.loads(Path(job["audit"]).read_text(encoding="utf-8"))
    # The workbook and the report must be hashed into the audit, so a later edit
    # to either can be detected.
    assert audit["output_workbook"]["sha256"] == sha256(job["workbook"])
    assert audit["output_report"]["sha256"] == sha256(job["report"])
    # Scores may only be read after native Excel recalculation.
    authority = audit["calculation_authority"]["recalculation"]
    assert authority["engine"] == "Microsoft Excel", (
        "scores were not read from a native Excel recalculation")


def test_boundary_values_score_as_the_charter_states(service, tmp_path):
    """Project-specific: the charter's own tier boundaries.

    Add examples/boundary_cases.json as

        {"cases": [{"facts": {...}, "expected": {"<output_id>": 100}}, ...]}

    then this test checks each boundary against the charter's stated tier. An
    unattested formula is an untested formula, and the boundaries are exactly
    where an off-by-one in a derived formula hides.
    """
    path = PROJECT / "examples" / "boundary_cases.json"
    if not path.is_file():
        pytest.skip("no examples/boundary_cases.json: add the charter's tier boundaries")
    cases = json.loads(path.read_text(encoding="utf-8"))["cases"]
    assert cases, "boundary_cases.json declares no cases"
    for index, case in enumerate(cases):
        result = service.run_job_from_facts(case["facts"], tmp_path / f"case{index}")
        for output_id, expected in case["expected"].items():
            assert result["scores"][output_id] == expected, (
                f"case {index} ({case['facts']}): {output_id} was "
                f"{result['scores'][output_id]}, charter states {expected}")
