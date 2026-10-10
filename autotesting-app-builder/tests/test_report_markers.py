"""Marker inventory for the blank DOCX report template.

`scripts/report_markers.py` is Skill-only: it does not exist in the upstream
AutoTesting repository, so these tests are new rather than ported. It exists
because the vendored inspector cannot see a bare `{{ }}` slot — its pattern
requires at least one identifier character between the braces.
"""
from __future__ import annotations
import sys, tempfile, unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from docx import Document
import report_markers


class MarkerInventoryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def build(self, paragraphs=(), header=None, name="template.docx"):
        path = self.root / name
        document = Document()
        for text in paragraphs:
            document.add_paragraph(text)
        if header is not None:
            document.sections[0].header.paragraphs[0].text = header
        document.save(path)
        return path

    def build_split(self, text_a, text_b, name="split.docx"):
        """A marker whose braces are split across two runs by a formatting change."""
        path = self.root / name
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run(text_a)
        paragraph.add_run(text_b)
        document.save(path)
        return path

    def test_bare_slots_are_detected_and_located(self):
        report = report_markers.extract_markers(self.build(["编号：{{}}", "得分：{{  }}"]))
        self.assertEqual(len(report["slots"]), 2)
        self.assertEqual(report["names"], [])
        self.assertTrue(all(item["locator"].startswith("word/") for item in report["slots"]))

    def test_named_markers_are_deduplicated_and_sorted(self):
        report = report_markers.extract_markers(self.build(["{{ score }}", "{{ alpha }}", "{{ score }}"]))
        self.assertEqual(report["names"], ["alpha", "score"])
        self.assertEqual(report["slots"], [])

    def test_marker_split_across_runs_is_still_found(self):
        # Word breaks a sentence across runs at every formatting change, so a
        # marker can span several w:t nodes. Reading run-by-run would miss it.
        report = report_markers.extract_markers(self.build_split("编号：{{ sc", "ore }}"))
        self.assertEqual(report["names"], ["score"])

    def test_bare_slot_split_across_runs_is_still_found(self):
        report = report_markers.extract_markers(self.build_split("编号：{{", "}}"))
        self.assertEqual(len(report["slots"]), 1)

    def test_unsupported_expression_is_surfaced_not_ignored(self):
        # A dotted name renders under docxtpl but is invisible to every
        # Skill-side walker, so it must be reported rather than passed silently.
        report = report_markers.extract_markers(self.build(["{{ vehicle.name }}"]))
        self.assertEqual(report["names"], [])
        self.assertEqual(len(report["unsupported"]), 1)
        self.assertIn("vehicle.name", report["unsupported"][0]["expression"])

    def test_all_three_bracketed_families_are_reported(self):
        report = report_markers.extract_markers(self.build(["[[square]]", "<<angle>>", "【corner】"]))
        self.assertEqual({item["family"] for item in report["bracketed"]},
                         {"square", "angle", "corner"})

    def test_corner_brackets_do_not_block_but_others_do(self):
        # 【 】 is ordinary Chinese punctuation; [[ ]] and << >> are deliberate
        # placeholder syntax that must be converted to {{ }}.
        corner = report_markers.extract_markers(self.build(["【注意事项】"]))
        self.assertEqual(report_markers.blocking_bracketed(corner), [])

        square = report_markers.extract_markers(self.build(["[[report_id]]"]))
        self.assertEqual([item["family"] for item in report_markers.blocking_bracketed(square)], ["square"])

    def test_control_syntax_is_detected(self):
        report = report_markers.extract_markers(self.build(["{% if x %}"]))
        self.assertTrue(report["has_control_syntax"])

    def test_header_markers_are_included(self):
        report = report_markers.extract_markers(self.build(["body {{ a }}"], header="页眉 {{ b }}"))
        self.assertEqual(report["names"], ["a", "b"])

    def test_source_sha256_matches_the_file(self):
        path = self.build(["{{ score }}"])
        self.assertEqual(report_markers.extract_markers(path)["source_sha256"],
                         report_markers.sha256_file(path))

    def test_template_field_names_is_the_named_set(self):
        path = self.build(["{{ score }}", "{{}}"])
        self.assertEqual(report_markers.template_field_names(path), frozenset({"score"}))

    def test_template_error_rejects_a_template_with_no_markers(self):
        # A DOCX with no markers is a completed report, not a blank template.
        report = report_markers.extract_markers(self.build(["已完成报告，无标记"]))
        self.assertFalse(report_markers.has_fill_markers(report))
        self.assertIn("不是空白模板", report_markers.template_error(report))

    def test_template_error_rejects_blocking_brackets_alongside_real_markers(self):
        # The bracketed check only applies once the template is known to be blank:
        # a document with no {{ }} at all is reported as "not a blank template"
        # first, which takes precedence over the mixed-marker complaint.
        report = report_markers.extract_markers(self.build(["得分：{{ score }}", "编号：[[report_id]]"]))
        self.assertIn("非 {{}} 的占位符标记", report_markers.template_error(report))

    def test_template_error_rejects_control_syntax(self):
        report = report_markers.extract_markers(self.build(["得分：{{ score }}", "{% if x %}"]))
        self.assertIn("控制语法", report_markers.template_error(report))

    def test_template_error_is_none_for_a_clean_named_template(self):
        report = report_markers.extract_markers(self.build(["得分：{{ score }}", "【说明】"]))
        self.assertIsNone(report_markers.template_error(report))

    def test_summarize_counts_slots_and_lists_locations(self):
        summary = report_markers.summarize(report_markers.extract_markers(self.build(["{{}}", "{{}}"])))
        self.assertEqual(summary["slots"], 2)
        self.assertEqual(len(summary["slotLocations"]), 2)


if __name__ == "__main__":
    unittest.main()
