from __future__ import annotations
import json
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import Workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from init_project import initialize
from test_three_document_flow import make_pdf


class GenericProjectInitializationTest(unittest.TestCase):
    def test_three_new_inputs_create_review_workspace_without_business_guesses(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);pdf=root/'policy.pdf';make_pdf(pdf)
            book=Workbook();book.active.title='Observations';book.active['A2']=5;book.save(root/'measurements.xlsx');book.close()
            doc=Document();doc.add_table(rows=1,cols=2);doc.tables[0].cell(0,0).text='Score';doc.tables[0].cell(0,1).text='{{}}';doc.save(root/'blank.docx')
            original=[item.read_bytes() for item in (pdf,root/'measurements.xlsx',root/'blank.docx')]
            output=root/'new-project'
            result=initialize(pdf,root/'measurements.xlsx',root/'blank.docx',output,project_name='Unrelated domain')
            self.assertEqual(result['status'],'review_required')
            self.assertEqual(result['report_positions'],1)
            self.assertEqual([item.read_bytes() for item in (pdf,root/'measurements.xlsx',root/'blank.docx')],original)
            self.assertEqual(json.loads((output/'review/rule-plan.json').read_text())['rules'],[])
            scope=json.loads((output/'review/rule-plan.json').read_text())['scope']
            self.assertEqual(scope['type'],'workbook')
            self.assertEqual(scope['score_targets'],[])
            pdf_inventory=json.loads((output/'analysis/pdf.json').read_text())
            self.assertEqual(pdf_inventory['schema_version'],'gamma-charter-source-v1')
            self.assertEqual(pdf_inventory['pages'],[{'page_index':1}])
            self.assertEqual(json.loads((output/'review/excel-map-plan.json').read_text())['output_cells'],{})
            report_position=json.loads((output/'review/report-map-plan.json').read_text())['positions'][0]
            self.assertIsNone(report_position['source'])
            self.assertEqual(report_position['context']['row_text'],'Score | {{}}')
            self.assertEqual(report_position['context']['table_column'],2)
            self.assertTrue((output/'src/app_service.py').is_file())
            with self.assertRaisesRegex(ValueError,'already exists'):
                initialize(pdf,root/'measurements.xlsx',root/'blank.docx',output,project_name='Again')

    def test_markerless_report_does_not_create_project(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);pdf=root/'policy.pdf';make_pdf(pdf)
            book=Workbook();book.save(root/'score.xlsx');book.close()
            doc=Document();doc.add_paragraph('Completed report');doc.save(root/'complete.docx')
            output=root/'blocked'
            with self.assertRaisesRegex(ValueError,'未发现任何'):
                initialize(pdf,root/'score.xlsx',root/'complete.docx',output,project_name='Blocked')
            self.assertFalse(output.exists())


if __name__=='__main__':unittest.main()
