"""Constructed engineering fixtures: formula authoring is never business approval."""
from __future__ import annotations

import json
from pathlib import Path
import sys
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import write_scoring_formulas as writer


class FormulaWritebackTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / "score.xlsx"
        self.plan_path = self.root / "plan.json"
        self.staged = self.root / "staged.xlsx"
        self.manifest = self.root / "stage.json"
        self.evidence = self.root / "native.json"
        self.receipt = self.root / "receipt.json"
        self.scoring_spec = self.root / "评分细则.md"
        self.rules = self.root / "rules.json"

    def fixture(self, mode):
        book = Workbook(); ws = book.active; ws.title = "评分"
        ws.merge_cells("A1:C1"); ws["A1"] = "工程构造评分表"
        ws["A2"] = "输入"; ws["B2"] = 0; ws["C2"] = None
        ws["A3"] = "第二项"; ws["B3"] = 0; ws["C3"] = None
        ws["C2"].fill = PatternFill("solid", fgColor="FFFF00")
        ws["C3"].fill = PatternFill("solid", fgColor="FFFF00")
        ws["C4"] = "=SUM(C2:C3)" if mode in ("aggregate", "partial") else None
        if mode == "partial": ws["C2"] = "=IF(B2<=10,2,0)"
        validation = DataValidation(type="whole", operator="between", formula1="0", formula2="20")
        ws.add_data_validation(validation); validation.add(ws["B2"])
        chart = BarChart(); chart.add_data(Reference(ws, min_col=2, min_row=2, max_row=3)); ws.add_chart(chart, "E2")
        book.save(self.source); book.close()
        original = writer.sha(self.source)
        formulas = []
        for row in ((3,) if mode == "partial" else (2, 3)):
            formulas.append({"sheet": "评分", "cell": f"C{row}", "original": None,
                             "formula": f"=IF(B{row}<=10,2,0)", "inputs": [f"评分!B{row}"],
                             "pdf_refs": [{"page": 1, "location": f"rule-{row}"}],
                             "rule_explanation": "输入不超过 10 得 2 分", "boundary_checks": [10, 11], "risks": [],
                             "rule_ids": ["R1"], "review_status": "confirmed", "reviewer": "Fixture reviewer",
                             "reviewed_at": "2026-10-10T12:00:00+08:00"})
        self.scoring_spec.write_text('# 构造评分细则\n\nPDF: '+ 'a'*64 +'\n\nR1',encoding='utf-8')
        self.rules.write_text(json.dumps({'source_sha256':'a'*64,'release_blocked':False,
                                          'rules':[{'id':'R1','status':'confirmed'}]}),encoding='utf-8')
        expected = [{"sheet": "评分", "cell": f"C{r}", "value": 2} for r in ((3,) if mode == "partial" else (2, 3))]
        self.plan = {"workbook": str(self.source), "source_sha256": original, "formulas": formulas,
                     "review":{"scoring_spec":{"path":str(self.scoring_spec),"sha256":writer.sha(self.scoring_spec)},
                               "rules":{"path":str(self.rules),"sha256":writer.sha(self.rules)}},
                     "boundary_cases": [{"inputs": [{"sheet": "评分", "cell": "B2", "value": 10},
                                                    {"sheet": "评分", "cell": "B3", "value": 10}], "expected": expected},
                                        {"inputs": [{"sheet": "评分", "cell": "B2", "value": 11},
                                                    {"sheet": "评分", "cell": "B3", "value": 11}],
                                         "expected": [{**item, "value": 0} for item in expected]}]}
        self.plan_path.write_text(json.dumps(self.plan, ensure_ascii=False), encoding="utf-8")
        return original

    def fake_excel(self, *args, **kwargs):
        manifest = json.loads(self.manifest.read_text())
        evidence = {"status": "passed", "engine": "Microsoft Excel", "excel_version": "TEST_DOUBLE_ONLY",
                    "plan_sha256": manifest["plan_sha256"], "source_sha256": manifest["source_sha256"],
                    "staged_sha256": manifest["staged_sha256"],
                    "cases": [{"status": "passed"} for _ in self.plan["boundary_cases"]]}
        self.evidence.write_text(json.dumps(evidence), encoding="utf-8")
        return SimpleNamespace(returncode=0, stderr="", stdout="")

    def test_missing_aggregate_and_partial_workbooks_write_original_path(self):
        for mode in ("empty", "aggregate", "partial"):
            with self.subTest(mode=mode):
                self.source.unlink(missing_ok=True)
                self.staged.unlink(missing_ok=True); self.manifest.unlink(missing_ok=True)
                self.evidence.unlink(missing_ok=True); self.receipt.unlink(missing_ok=True)
                original = self.fixture(mode)
                writer.prepare(self.plan_path, self.staged, self.manifest)
                with patch.object(writer.subprocess, "run", side_effect=self.fake_excel):
                    result = writer.commit(self.plan_path, self.manifest, self.evidence, self.receipt)
                self.assertEqual(result["original_sha256"], original)
                self.assertEqual(writer.sha(Path(result["backup_path"])), original)
                self.assertEqual(writer.sha(self.source), result["new_sha256"])
                with ZipFile(self.source) as changed, ZipFile(result["backup_path"]) as before:
                    changed_parts = {n for n in before.namelist() if before.read(n) != changed.read(n)}
                    self.assertEqual(changed_parts, {"xl/worksheets/sheet1.xml"})
                book = load_workbook(self.source)
                try:
                    ws = book["评分"]
                    for formula in self.plan["formulas"]:
                        self.assertEqual(ws[formula["cell"]].value, formula["formula"])
                    self.assertEqual(len(ws._charts), 1)
                    self.assertEqual(len(ws.data_validations.dataValidation), 1)
                    self.assertIn("A1:C1", [str(x) for x in ws.merged_cells.ranges])
                    self.assertEqual(ws["C2"].fill.fgColor.rgb, "00FFFF00")
                    if mode in ("aggregate", "partial"):
                        self.assertEqual(ws["C4"].value, "=SUM(C2:C3)")
                finally:
                    book.close()

    def test_prepare_writes_unformatted_blank_cell_and_missing_row(self):
        for target in ('B2','C5'):
            with self.subTest(target=target):
                self.source.unlink(missing_ok=True);self.staged.unlink(missing_ok=True);self.manifest.unlink(missing_ok=True)
                book=Workbook();ws=book.active;ws.title='评分';ws['A2']=5;book.save(self.source);book.close()
                formula='=IF(A2<=10,2,0)'
                plan={'workbook':'score.xlsx','source_sha256':writer.sha(self.source),
                      'formulas':[{'sheet':'评分','cell':target,'original':None,'formula':formula,
                                   'inputs':['评分!A2'],'pdf_refs':[{'page':1,'location':'rule'}],
                                   'rule_explanation':'constructed','boundary_checks':[10,11],'risks':[]}],
                      'boundary_cases':[{'inputs':[{'sheet':'评分','cell':'A2','value':10}],
                                         'expected':[{'sheet':'评分','cell':target,'value':2}]}]}
                self.plan_path.write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
                original=writer.sha(self.source)
                writer.prepare(self.plan_path,self.staged,self.manifest)
                self.assertEqual(writer.sha(self.source),original)
                self.assertEqual(load_workbook(self.staged)['评分'][target].value,formula)

    def test_rejects_stale_plan_and_stage_without_touching_source(self):
        original = self.fixture("empty")
        writer.prepare(self.plan_path, self.staged, self.manifest)
        self.plan["formulas"][0]["formula"] = "=IF(B2<=10,0,0)"
        self.plan_path.write_text(json.dumps(self.plan, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "staged formula mismatch|stage manifest"):
            writer.commit(self.plan_path, self.manifest, self.evidence, self.receipt)
        self.assertEqual(writer.sha(self.source), original)
        self.assertFalse(self.receipt.exists())

    def test_native_failure_keeps_original_and_no_backup(self):
        original = self.fixture("empty")
        writer.prepare(self.plan_path, self.staged, self.manifest)
        with patch.object(writer.subprocess, "run", return_value=SimpleNamespace(returncode=2, stderr="Excel failed", stdout="")):
            with self.assertRaisesRegex(ValueError, "native validation failed"):
                writer.commit(self.plan_path, self.manifest, self.evidence, self.receipt)
        self.assertEqual(writer.sha(self.source), original)
        self.assertFalse(list(self.root.glob("*.backup-*")))

    def test_pending_or_changed_spec_blocks_commit_before_excel(self):
        original=self.fixture("empty")
        writer.prepare(self.plan_path,self.staged,self.manifest)
        self.plan['formulas'][0]['review_status']='pending'
        self.plan_path.write_text(json.dumps(self.plan,ensure_ascii=False),encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'named confirmation'):
            writer.commit(self.plan_path,self.manifest,self.evidence,self.receipt)
        self.assertEqual(writer.sha(self.source),original)
        self.assertFalse(self.evidence.exists())
        self.plan['formulas'][0]['review_status']='confirmed'
        self.plan_path.write_text(json.dumps(self.plan,ensure_ascii=False),encoding='utf-8')
        self.scoring_spec.write_text(self.scoring_spec.read_text(encoding='utf-8')+'\nchanged',encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'scoring_spec file changed'):
            writer.commit(self.plan_path,self.manifest,self.evidence,self.receipt)

    def test_cycle_and_occupied_cell_block_prepare(self):
        original = self.fixture("empty")
        self.plan["formulas"][0]["formula"] = "=IF(B2<=10,C3,0)"
        self.plan["formulas"][1]["formula"] = "=IF(B3<=10,C2,0)"
        self.plan_path.write_text(json.dumps(self.plan, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "formula cycle"):
            writer.prepare(self.plan_path, self.staged, self.manifest)
        self.assertEqual(writer.sha(self.source), original)
        self.assertFalse(self.staged.exists())

    def test_missing_input_and_cycle_through_existing_aggregate_block(self):
        original = self.fixture("aggregate")
        self.plan["formulas"][0]["formula"] = "=IF(B2<=10,C4,0)"
        self.plan_path.write_text(json.dumps(self.plan, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "formula cycle"):
            writer.prepare(self.plan_path, self.staged, self.manifest)
        self.plan["formulas"][0]["formula"] = "=IF(B2<=10,B5,0)"
        self.plan_path.write_text(json.dumps(self.plan, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "missing formula input declaration"):
            writer.prepare(self.plan_path, self.staged, self.manifest)
        self.assertEqual(writer.sha(self.source), original)

    def test_post_replace_failure_restores_source_and_preserves_backup(self):
        original = self.fixture("empty")
        writer.prepare(self.plan_path, self.staged, self.manifest)
        real_load = writer.load_workbook
        calls = 0
        def fail_after_replace(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls >= 3:
                raise OSError("simulated reopen failure")
            return real_load(*args, **kwargs)
        with patch.object(writer.subprocess, "run", side_effect=self.fake_excel), patch.object(writer, "load_workbook", side_effect=fail_after_replace):
            with self.assertRaisesRegex(OSError, "simulated reopen"):
                writer.commit(self.plan_path, self.manifest, self.evidence, self.receipt)
        self.assertEqual(writer.sha(self.source), original)
        backups = list(self.root.glob("*.backup-*"))
        self.assertEqual(len(backups), 1)
        self.assertEqual(writer.sha(backups[0]), original)
        self.assertFalse(self.receipt.exists())


if __name__ == "__main__": unittest.main()
