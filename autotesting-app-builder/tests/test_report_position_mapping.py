"""Position-by-position DOCX mapping and actual render, with constructed data."""
from __future__ import annotations

import json
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from docxtpl import DocxTemplate

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import map_report_template, report_markers


class ReportPositionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / "blank.docx"
        self.output = self.root / "named.docx"
        self.manifest = self.root / "mapping.json"
        self.plan_path = self.root / "plan.json"

    def fixture(self):
        doc = Document()
        doc.add_paragraph("静态说明：保持原文。")
        p = doc.add_paragraph("样车：")
        p.add_run("{{"); p.add_run("}}")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "评分"
        table.cell(0, 1).text = "{{ }}"
        table.cell(1, 0).merge(table.cell(1, 1)).text = "备注：{{ note }}"
        doc.sections[0].header.paragraphs[0].text = "页眉编号：{{}}"
        doc.sections[0].footer.paragraphs[0].text = "页脚日期：{{ }}"
        doc.save(self.source)
        inventory = report_markers.extract_markers(self.source)
        def assign(pos):
            context = pos["context"]
            if "样车" in context: return "vehicle", {"kind": "input", "id": "vehicle"}
            if "评分" in context or pos["part"] == "word/document.xml" and "{{ }}" in context: return "score", {"kind": "score", "id": "score"}
            if "备注" in context: return "note", {"kind": "input", "id": "note"}
            if "页眉" in context: return "audit_id", {"kind": "system", "id": "audit_id"}
            if "页脚" in context: return "date", {"kind": "input", "id": "date"}
            raise AssertionError(context)
        positions = []
        for pos in inventory["positions"]:
            field, source = assign(pos)
            positions.append({"position_id": pos["position_id"], "field_id": field, "source": source})
        plan = {"template": str(self.source), "output": str(self.output), "manifest": str(self.manifest),
                "source_sha256": inventory["source_sha256"], "positions": positions,
                "available_sources": {"input": ["vehicle", "note", "date"], "score": ["score"], "system": ["audit_id"]}}
        self.plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        return plan

    def test_all_positions_named_and_rendered(self):
        plan = self.fixture(); before = report_markers.sha256_file(self.source)
        result = map_report_template.map_template(self.plan_path)
        self.assertEqual(len(result["positions"]), 5)
        self.assertEqual(before, report_markers.sha256_file(self.source))
        self.assertEqual(set(report_markers.extract_markers(self.output)["names"]), {"vehicle", "score", "note", "audit_id", "date"})
        rendered = self.root / "rendered.docx"
        doc = DocxTemplate(str(self.output))
        doc.render({"vehicle": "工程样车", "score": 2, "note": "人工录入", "audit_id": "X001", "date": "2026-10-10"})
        doc.save(rendered)
        inventory = report_markers.extract_markers(rendered)
        self.assertEqual(inventory["positions"], [])
        view = Document(rendered)
        self.assertIn("静态说明：保持原文。", [p.text for p in view.paragraphs])
        self.assertIn("工程样车", " ".join(p.text for p in view.paragraphs))
        self.assertEqual(view.tables[0].cell(0, 1).text, "2")
        self.assertIn("人工录入", view.tables[0].cell(1, 0).text)
        self.assertIn("X001", view.sections[0].header.paragraphs[0].text)
        self.assertIn("2026-10-10", view.sections[0].footer.paragraphs[0].text)

    def test_missing_or_unknown_source_blocks_without_output(self):
        plan = self.fixture()
        plan["positions"].pop()
        self.plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "every marker position"):
            map_report_template.map_template(self.plan_path)
        self.assertFalse(self.output.exists())
        plan = self.fixture(); plan["positions"][0]["source"] = {"kind": "score", "id": "unknown"}
        self.plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "unresolved data source"):
            map_report_template.map_template(self.plan_path)
        self.assertFalse(self.output.exists())

    def test_markerless_document_blocks(self):
        Document().save(self.source)
        plan = {"template": str(self.source), "output": str(self.output), "manifest": str(self.manifest),
                "source_sha256": report_markers.sha256_file(self.source), "positions": [], "available_sources": {}}
        self.plan_path.write_text(json.dumps(plan), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "未发现任何"):
            map_report_template.map_template(self.plan_path)


if __name__ == "__main__": unittest.main()
