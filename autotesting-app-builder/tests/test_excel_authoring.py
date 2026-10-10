"""Read-only assessment workbook analysis contracts.

Ported from the upstream AutoTesting repository's tests/test_excel_authoring.py.
Exercises the frozen vendored `ivista_engine.authoring.excel_analyzer`, so it also
pins the vendored contract: an accidental edit to vendor/ fails these assertions.
The only change from upstream is the sys.path bootstrap below.
"""
from __future__ import annotations
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "vendor"))

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from ivista_engine.authoring.errors import AuthoringError
from ivista_engine.authoring.excel_analyzer import analyze_workbook, write_inventory

class ExcelAnalyzerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.workbook_path = self.root / "人工评分空表.xlsx"
        workbook = Workbook()
        summary = workbook.active
        summary.title = "评分汇总"
        summary.merge_cells("A1:C1")
        summary["A1"] = "封闭场地评分结果"
        summary["A3"] = "场景"
        summary["B3"] = "实测结果"
        summary["C3"] = "得分"
        summary["A4"] = "隧道测试"
        summary["B4"] = None
        summary["C4"] = '=IF(B4="通过",2,0)'
        summary["B4"].number_format = "@"
        summary["C4"].number_format = "0.0"
        validation = DataValidation(type="list", formula1='"通过,未通过"', allow_blank=True)
        summary.add_data_validation(validation)
        validation.add(summary["B4"])
        detail = workbook.create_sheet("细项录入")
        detail["A1"] = "场景"
        detail["B1"] = "成功次数"
        detail["C1"] = "总次数"
        detail["A2"] = "隧道测试"
        detail["B2"] = 0
        detail["C2"] = 10
        detail.sheet_state = "hidden"
        from openpyxl.workbook.defined_name import DefinedName
        workbook.defined_names.add(DefinedName("TestInput", attr_text="'评分汇总'!$B$4"))
        workbook.save(self.workbook_path)
        workbook.close()
        self.before_hash = hashlib.sha256(self.workbook_path.read_bytes()).hexdigest()

    def tearDown(self):
        self.temp.cleanup()

    def test_inventory_captures_structure_without_modifying_source(self):
        inventory = analyze_workbook(self.workbook_path)
        self.assertEqual(inventory["source_sha256"], self.before_hash)
        self.assertFalse(inventory["write_supported"])
        self.assertEqual(hashlib.sha256(self.workbook_path.read_bytes()).hexdigest(), self.before_hash)
        self.assertIn("A1:C1", inventory["sheets"][0]["merged_ranges"])
        self.assertEqual(inventory["sheets"][1]["state"], "hidden")
        self.assertTrue(inventory["defined_names"])
        cells = {(item["sheet"], item["cell"]): item for item in inventory["cells"]}
        self.assertTrue(cells[("评分汇总", "B4")]["validations"])
        self.assertEqual(cells[("评分汇总", "C4")]["value_type"], "formula")
        self.assertIn({"sheet": None, "cell": "B4"}, cells[("评分汇总", "C4")]["formula_references"])

    def test_inventory_json_and_invalid_suffix(self):
        path = self.root / "workbook_inventory.json"
        first = write_inventory(self.workbook_path, path)
        second = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(first["source_sha256"], second["source_sha256"])
        with self.assertRaises(AuthoringError) as caught:
            analyze_workbook(self.root / "fake.xlsm")
        self.assertEqual(caught.exception.code, "EXCEL_INPUT_INVALID")

if __name__ == "__main__":
    unittest.main()
