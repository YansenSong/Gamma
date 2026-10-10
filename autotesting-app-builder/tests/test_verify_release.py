"""Release-gate contract of scripts/verify_release.py and scripts/write_manifest.py.

Ports the upstream AutoTesting repository's tests/test_release_verification.py
and extends it. Upstream wrote that file for pytest (`tmp_path` fixture); this
version is stdlib unittest so the Skill keeps one test command and no test
dependency. Two adaptations were required:

  1. Paths resolve inside the Skill folder instead of the repository checkout.
  2. Upstream's fixture stubbed its sources with raw bytes and declared
     `report_output.fields: {}`. The Skill's validator now parses the report
     template and requires a non-empty field list whose names equal the
     template's {{ }} markers, so the fixture builds a real DOCX and a real
     workbook with a formula. `create_manifest` calls the validator, so without
     this the fixture could not produce a manifest at all.

The added case is `test_empty_tests_directory_is_rejected`: upstream's own
fixture created an empty `tests/` directory and passed `tests_status="passed"`,
so the hole was never exercised.
"""
from __future__ import annotations
import hashlib, json, subprocess, sys, tempfile, unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from docx import Document
from openpyxl import Workbook
from write_manifest import create_manifest

VERIFY = SCRIPTS / "verify_release.py"


def sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class ReleaseVerificationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / "sample"
        self.root.mkdir()

    def make_project(self, *, with_test_module=True, release_status="engineering_sample"):
        root = self.root
        for folder in ("config", "inputs", "src", "tests"):
            (root / folder).mkdir(parents=True, exist_ok=True)
        (root / "README.md").write_text("sample", encoding="utf-8")
        (root / "build.ps1").write_text("", encoding="utf-8")
        (root / "src" / "main.py").write_text("", encoding="utf-8")
        if with_test_module:
            (root / "tests" / "test_sample.py").write_text("def test_ok():\n    assert True\n",
                                                           encoding="utf-8")
        (root / "inputs" / "charter.pdf").write_bytes(b"charter")
        document = Document(); document.add_paragraph("得分：{{ score }}")
        document.save(root / "inputs" / "report.docx")
        book = Workbook(); book.active.title = "Inputs"; book["Inputs"]["B4"] = "=IF(B2<=200,100,0)"
        book.save(root / "inputs" / "score.xlsx"); book.close()
        sources = [{"role": role, "path": f"inputs/{name}", "sha256": sha256(root / "inputs" / name)}
                   for role, name in (("charter", "charter.pdf"), ("workbook", "score.xlsx"),
                                      ("report_template", "report.docx"))]
        spec = {"schema_version": "autotesting-project-v1",
                "application": {"id": "sample", "name": "Sample", "version": "0.1.0"},
                "release_status": release_status, "sources": sources, "business_decisions": [],
                "scoring": {"mode": "excel_native",
                            "inputs": [{"id": "latency", "type": "number", "required": True}],
                            "formula_bindings": [{"output": "score", "sheet": "Inputs", "cell": "B4",
                                                  "formula": "=IF(B2<=200,100,0)",
                                                  "source_refs": ["p1"]}]},
                "excel_output": {"template": "inputs/score.xlsx",
                                 "input_cells": {"latency": {"sheet": "Inputs", "cell": "B2"}},
                                 "output_cells": {"score": {"sheet": "Inputs", "cell": "B4"}}},
                "report_output": {"template": "inputs/report.docx",
                                  "fields": {"score": {"kind": "score", "id": "score"}}}}
        (root / "config" / "project.json").write_text(json.dumps(spec), encoding="utf-8")
        manifest = create_manifest(root, tests_status="passed", tests_command="pytest",
                                   build_status="not_run", smoke_status="not_run",
                                   executable=None, limitations=[])
        (root / "delivery-manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
        return root

    def verify(self):
        return subprocess.run([sys.executable, str(VERIFY), str(self.root)],
                              capture_output=True, text=True)

    def test_a_complete_project_passes(self):
        self.make_project()
        result = self.verify()
        self.assertEqual(result.returncode, 0, result.stdout)

    def test_manifest_only_release_promotion_is_rejected(self):
        self.make_project()
        path = self.root / "delivery-manifest.json"
        manifest = json.loads(path.read_text(encoding="utf-8"))
        manifest["release_status"] = "business_validated"
        path.write_text(json.dumps(manifest), encoding="utf-8")
        result = self.verify()
        self.assertEqual(result.returncode, 1)
        self.assertIn("release_status differs from project.json", result.stdout)

    def test_project_spec_mutation_after_manifest_is_rejected(self):
        self.make_project()
        path = self.root / "config" / "project.json"
        spec = json.loads(path.read_text(encoding="utf-8"))
        spec["application"]["version"] = "9.9.9"
        path.write_text(json.dumps(spec), encoding="utf-8")
        result = self.verify()
        self.assertEqual(result.returncode, 1)
        self.assertIn("project_spec_sha256 is stale", result.stdout)

    def test_empty_tests_directory_is_rejected(self):
        # `mkdir tests` plus a self-reported tests_status="passed" used to satisfy
        # the gate. The directory must carry a real test module.
        self.make_project(with_test_module=False)
        result = self.verify()
        self.assertEqual(result.returncode, 1)
        self.assertIn("tests/test_*.py", result.stdout)

    def test_missing_tests_directory_is_rejected(self):
        self.make_project()
        for child in (self.root / "tests").iterdir():
            child.unlink()
        (self.root / "tests").rmdir()
        result = self.verify()
        self.assertEqual(result.returncode, 1)
        self.assertIn("tests", result.stdout)

    def test_tampered_artifact_is_rejected(self):
        self.make_project()
        (self.root / "src" / "main.py").write_text("tampered", encoding="utf-8")
        result = self.verify()
        self.assertEqual(result.returncode, 1)
        self.assertIn("sha256 mismatch", result.stdout)


if __name__ == "__main__":
    unittest.main()
