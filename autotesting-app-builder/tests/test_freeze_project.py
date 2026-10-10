from __future__ import annotations
import json
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import Workbook,load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from init_project import initialize
from map_report_template import map_template
from publish_scoring_rules import publish
from freeze_project import freeze
from test_three_document_flow import make_pdf


def save(path,value):path.write_text(json.dumps(value,ensure_ascii=False),encoding='utf-8')


class FreezeGenericProjectTest(unittest.TestCase):
    def test_reviewed_artifacts_freeze_without_domain_constants(self):
        with tempfile.TemporaryDirectory() as directory:
            base=Path(directory);pdf=base/'other.pdf';make_pdf(pdf)
            book=Workbook();book.active.title='Numbers';book.active['A2']=5;book.save(base/'table.xlsx');book.close()
            doc=Document();doc.add_paragraph('Result {{}}');doc.save(base/'report.docx')
            root=base/'project';initialize(pdf,base/'table.xlsx',base/'report.docx',root,project_name='Unrelated Project')
            inventory=json.loads((root/'analysis/pdf.json').read_text());ref=inventory['pages'][0]['blocks'][0]['ref']
            rule_plan={'source_sha256':inventory['source_sha256'],'reviewed_pages':[1],
                       'rules':[{'id':'R1','statement':'Synthetic score rule','status':'pending','source_refs':[ref]}]}
            save(root/'review/rule-plan.json',rule_plan)
            publish(root/'analysis/pdf.json',root/'review/rule-plan.json',root/'analysis/评分细则.md',root/'analysis/scoring_rules.json',root/'analysis/open_questions.json')
            workbook=load_workbook(root/'templates/score.xlsx');workbook['Numbers']['B2']='=IF(A2<=10,2,0)';workbook.save(root/'templates/score.xlsx');workbook.close()
            excel=json.loads((root/'review/excel-map-plan.json').read_text());excel['input_cells']={'observation':{'sheet':'Numbers','cell':'A2','review_status':'confirmed','source_refs':[ref]}}
            excel['output_cells']={'result':{'sheet':'Numbers','cell':'B2','review_status':'confirmed','source_refs':[ref]}};save(root/'review/excel-map-plan.json',excel)
            report_plan=json.loads((root/'review/report-map-plan.json').read_text());report_plan['available_sources']['score']=['result']
            report_plan['positions'][0].update({'field_id':'result','source':{'kind':'score','id':'result'},'review_status':'confirmed'})
            save(root/'review/report-map-plan.json',report_plan);map_template(root/'review/report-map-plan.json')
            project=json.loads((root/'review/project-plan.json').read_text());project['application']['id']='unrelated'
            project['formula_bindings']=[{'output':'result','sheet':'Numbers','cell':'B2','formula':'=IF(A2<=10,2,0)','source_refs':[ref]}]
            save(root/'review/project-plan.json',project)
            outcome=freeze(root);self.assertEqual(outcome['status'],'engineering_sample')
            self.assertTrue(outcome['business_review_blocked'])
            config=json.loads((root/'config/project.json').read_text())
            self.assertEqual(config['report_output']['fields']['result'],{'kind':'score','id':'result'})
            self.assertEqual(config['excel_output']['output_cells']['result']['cell'],'B2')
            with self.assertRaisesRegex(ValueError,'already configured'):freeze(root)


if __name__=='__main__':unittest.main()
