"""Three-file engineering flow. Excel COM is doubled here; Windows native gate is separate."""
from __future__ import annotations

import json
from pathlib import Path
import sys
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

import pdfplumber
from docx import Document
from docxtpl import DocxTemplate
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import map_report_template, report_markers, write_scoring_formulas as writer


def make_pdf(path):
    clause = b"Score 2 when input <= 10; otherwise 0."
    stream = b"BT /F1 12 Tf 72 720 Td (" + clause + b") Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    data = b"%PDF-1.4\n"; offsets = []
    for number, obj in enumerate(objects, 1):
        offsets.append(len(data)); data += str(number).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(data); data += b"xref\n0 6\n0000000000 65535 f \n"
    for offset in offsets: data += f"{offset:010d} 00000 n \n".encode()
    data += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF\n"
    path.write_bytes(data)


class ThreeDocumentFlow(unittest.TestCase):
    def test_constructed_pdf_blank_excel_marked_docx_to_report(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pdf = root / "rule.pdf"; make_pdf(pdf)
            with pdfplumber.open(pdf) as charter:
                clause = charter.pages[0].extract_text()
            self.assertIn("input <= 10", clause)
            xlsx = root / "score.xlsx"
            book = Workbook(); sheet = book.active; sheet.title = "Score"
            sheet["A1"] = "Input"; sheet["B1"] = "Score"; sheet["A2"] = 0
            sheet["B2"].number_format = "0"
            book.save(xlsx); book.close()
            original = writer.sha(xlsx)
            spec=root/'评分细则.md';rules=root/'rules.json'
            pdf_hash=writer.sha(pdf)
            spec.write_text('# 构造评分细则\n\nPDF: '+pdf_hash+'\n\nR1: '+clause,encoding='utf-8')
            rules.write_text(json.dumps({'source_sha256':pdf_hash,'release_blocked':False,
                                         'rules':[{'id':'R1','status':'confirmed'}]}),encoding='utf-8')
            plan = {"workbook": str(xlsx), "source_sha256": original,
                    "review":{"scoring_spec":{"path":str(spec),"sha256":writer.sha(spec)},
                              "rules":{"path":str(rules),"sha256":writer.sha(rules)}},
                    "formulas": [{"sheet": "Score", "cell": "B2", "original": None,
                                  "formula": "=IF(A2<=10,2,0)", "inputs": ["Score!A2"],
                                  "pdf_refs": [{"page": 1, "location": "text line 1"}],
                                  "rule_explanation": clause, "boundary_checks": [10, 11], "risks": [],
                                  "rule_ids":["R1"],"review_status":"confirmed","reviewer":"Fixture reviewer",
                                  "reviewed_at":"2026-10-10T12:00:00+08:00"}],
                    "boundary_cases": [{"inputs": [{"sheet": "Score", "cell": "A2", "value": value}],
                                        "expected": [{"sheet": "Score", "cell": "B2", "value": expected}]}
                                       for value, expected in ((10, 2), (11, 0))]}
            plan_path = root / "formula-plan.json"; plan_path.write_text(json.dumps(plan), encoding="utf-8")
            staged = root / "stage.xlsx"; manifest = root / "stage.json"
            writer.prepare(plan_path, staged, manifest)
            evidence = root / "native-evidence.json"; receipt = root / "receipt.json"
            def constructed_excel_gate(*args, **kwargs):
                stage = json.loads(manifest.read_text())
                evidence.write_text(json.dumps({"status": "passed", "engine": "Microsoft Excel",
                  "excel_version": "TEST_DOUBLE_ONLY", "plan_sha256": stage["plan_sha256"],
                  "source_sha256": stage["source_sha256"], "staged_sha256": stage["staged_sha256"],
                  "cases": [{"status": "passed"} for _ in plan["boundary_cases"]]}))
                return SimpleNamespace(returncode=0, stderr="", stdout="")
            with patch.object(writer.subprocess, "run", side_effect=constructed_excel_gate):
                written = writer.commit(plan_path, manifest, evidence, receipt)
            self.assertEqual(written["original_sha256"], original)
            self.assertEqual(load_workbook(xlsx)["Score"]["B2"].value, "=IF(A2<=10,2,0)")
            docx = root / "blank.docx"; doc = Document()
            doc.add_paragraph("Engineering example")
            doc.add_paragraph("Score: {{}}")
            doc.add_paragraph("Vehicle: {{ }}")
            doc.save(docx)
            positions = report_markers.extract_markers(docx)["positions"]
            mapping = {"template": str(docx), "output": str(root / "named.docx"),
                       "manifest": str(root / "report-map.json"), "source_sha256": report_markers.sha256_file(docx),
                       "available_sources": {"input": ["vehicle"], "score": ["score"], "system": []},
                       "positions": [{"position_id": p["position_id"], "field_id": "score" if "Score" in p["context"] else "vehicle",
                                      "source": {"kind": "score" if "Score" in p["context"] else "input",
                                                 "id": "score" if "Score" in p["context"] else "vehicle"}} for p in positions]}
            map_plan = root / "report-plan.json"; map_plan.write_text(json.dumps(mapping), encoding="utf-8")
            map_report_template.map_template(map_plan)
            template = DocxTemplate(mapping["output"]); template.render({"score": "2", "vehicle": "Constructed Car"})
            rendered = root / "rendered.docx"; template.save(rendered)
            result = Document(rendered)
            self.assertEqual([p.text for p in result.paragraphs], ["Engineering example", "Score: 2", "Vehicle: Constructed Car"])
            self.assertEqual(report_markers.extract_markers(rendered)["positions"], [])


if __name__ == "__main__": unittest.main()
