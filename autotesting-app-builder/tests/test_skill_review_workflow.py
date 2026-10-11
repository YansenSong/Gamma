"""Review workflow invariants using synthetic documents, never business approvals.

Ported from the upstream AutoTesting repository's tests/test_skill_review_workflow.py.

Two changes from upstream:
  1. ROOT/SCRIPTS/ASSETS resolve inside the Skill folder. Upstream reached into
     the checkout at `.agents/skills/autotesting-app-builder/`, which does not
     exist once the Skill ships standalone.
  2. The fixture declares `report_output.fields`. The Skill's validator now
     requires a non-empty field list whose names equal the template's {{ }}
     markers exactly; upstream's fixture predates that rule and declared `{}`,
     which the validator rejects before any review-workflow assertion runs.
"""
import csv
import importlib.util
import json
import sys
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from docx import Document
from openpyxl import Workbook

ROOT=Path(__file__).resolve().parents[1]
SCRIPTS=ROOT/'scripts'
ASSETS=ROOT/'assets/windows-app-template/src'
sys.path.insert(0,str(SCRIPTS));sys.path.insert(0,str(ASSETS));sys.path.insert(0,str(ROOT/'vendor'))

def load(name,path):
    spec=importlib.util.spec_from_file_location(name,path);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);return module

workflow=load('skill_review_workflow',SCRIPTS/'review_workflow.py')
inventory=load('skill_review_inventory',SCRIPTS/'build_review_inventory.py')
preview=load('skill_unscored_preview',ASSETS/'review_preview.py')


class ReviewWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp=TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
        for folder in ('config','analysis','templates','evidence/submissions'): (self.root/folder).mkdir(parents=True)
        document=Document();p=document.add_paragraph();p.add_run('{{ na');p.add_run('me }}');document.add_paragraph('Score: {{ score }}')
        document.save(self.root/'templates/report.docx')
        (self.root/'templates/charter.pdf').write_bytes(b'synthetic charter identity only')
        workbook=Workbook();workbook.active['A1']='=1+1';workbook.save(self.root/'templates/workbook.xlsx')
        sources=[{'role':role,'path':'templates/'+name,'sha256':workflow.file_hash(self.root/'templates'/name)} for role,name in
                 [('charter','charter.pdf'),('workbook','workbook.xlsx'),('report_template','report.docx')]]
        self.spec={'schema_version':'autotesting-project-v1','application':{'id':'synthetic-review','name':'Review','version':'1'},
          'release_status':'engineering_sample','sources':sources,'approvals':{},'business_decisions':[],
          'validation_state':{'status':'blocked','official_output_allowed':False,'reasons':['Synthetic only']},
          'scoring':{'mode':'excel_native','inputs':[{'id':'report_name','type':'string','required':False}],
                     'formula_bindings':[]},
          'excel_output':{'template':'templates/workbook.xlsx','input_cells':{},'output_cells':{}},
          'report_output':{'template':'templates/report.docx',
                           'fields':{'name':{'kind':'input','id':'report_name'},
                                     'score':{'kind':'score','id':'score'}}},
          'report_review':{'path':'config/report-field-review.json'},
          'review_preview':{'metadata_fields':{'name':{'label':'Name','input_id':'report_name'}}}}
        workflow.write(self.root/'config/project.json',self.spec)
        workflow.write(self.root/'analysis/report.json',{'body':[],'source_sha256':sources[-1]['sha256']})
        self.review=inventory.build(self.root);workflow.write(self.root/'config/report-field-review.json',self.review)
        self.destination=self.root/'analysis/review-package';self.package=workflow.export_package(self.root,self.destination)
        self.csv=self.destination/'opinions.csv';self.package_path=self.destination/'package.json'

    def rows(self):
        with self.csv.open(encoding='utf-8-sig',newline='') as stream:return list(csv.DictReader(stream))

    def save_rows(self,rows):
        with self.csv.open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.DictWriter(stream,fieldnames=list(rows[0]),lineterminator='\n');writer.writeheader();writer.writerows(rows)

    def check(self,previous=None):return workflow.revalidate(self.root,self.package_path,self.csv,previous)

    def submit(self,decision='accept',evidence='trial.txt'):
        path=self.root/'evidence/submissions'/evidence;path.write_text('synthetic record, not real trial',encoding='utf-8')
        rows=self.rows();row=rows[0];row.update(reviewer='Synthetic reviewer',reviewed_at='2026-10-10T12:00:00+08:00',
            decision=decision,reason='Synthetic test opinion',evidence_paths=json.dumps(['evidence/submissions/'+evidence]),
            evidence_sha256=json.dumps([workflow.file_hash(path)]),record_ids=json.dumps(['synthetic-record-1']))
        self.save_rows(rows);return path

    def test_inventory_explicit_metadata_and_unmapped_results(self):
        fields={x['field_id']:x for x in self.review['fields']}
        self.assertEqual(fields['name']['category'],'A');self.assertEqual(fields['score']['category'],'D')
        self.assertTrue(all(x['blocks_official'] for x in fields.values()))

    def test_empty_opinions_remain_pending(self):
        result=self.check();self.assertEqual(result['counts'],{'pending':len(self.package['items'])});self.assertFalse(result['official_output_allowed'])

    def test_accept_opinion_does_not_mutate_or_approve_project(self):
        before=(self.root/'config/project.json').read_bytes();self.submit();result=self.check()
        self.assertEqual(result['counts']['opinion_submitted'],1);self.assertFalse(result['official_output_allowed'])
        self.assertEqual(before,(self.root/'config/project.json').read_bytes());self.assertTrue(all(x['blocks_official'] for x in result['items']))

    def test_not_applicable_does_not_approve_empty_cache(self):
        self.submit('not_applicable');self.assertFalse(self.check()['official_output_allowed'])

    def test_stale_mapping_rejected(self):
        self.spec['report_output']['fields']['name']={'kind':'literal','value':'Changed'}
        workflow.write(self.root/'config/project.json',self.spec)
        with self.assertRaisesRegex(ValueError,'stale'):self.check()

    def test_source_mutation_rejected(self):
        (self.root/'templates/charter.pdf').write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError,'source hash mismatch'):self.check()

    def test_package_tampering_rejected(self):
        data=workflow.read(self.package_path);data['items'][0]['question']='changed';workflow.write(self.package_path,data)
        with self.assertRaisesRegex(ValueError,'content hash mismatch'):self.check()

    def test_unknown_and_duplicate_opinions_rejected(self):
        rows=self.rows();rows.append(dict(rows[0]));self.save_rows(rows)
        with self.assertRaisesRegex(ValueError,'duplicate'):self.check()
        rows=self.rows()[:-1];rows[0]['item_id']='unknown';self.save_rows(rows)
        with self.assertRaisesRegex(ValueError,'Unknown'):self.check()

    def test_opinion_subject_hash_rejected(self):
        rows=self.rows();rows[0]['content_sha256']='0'*64;self.save_rows(rows)
        with self.assertRaisesRegex(ValueError,'binding mismatch'):self.check()

    def test_missing_fields_and_timezone_block_submission(self):
        self.submit();rows=self.rows();rows[0].update(reviewer='',reviewed_at='2026-10-10T12:00:00');self.save_rows(rows)
        item=self.check()['items'][0];self.assertEqual(item['status'],'invalid_submission')
        self.assertIn('review_time_requires_timezone',item['errors']);self.assertIn('missing_reviewer',item['errors'])

    def test_evidence_mutation_is_detected_even_with_previous_record(self):
        path=self.submit();first=self.check();prior=self.root/'prior.json';workflow.write(prior,first)
        path.write_text('changed',encoding='utf-8');result=self.check(prior);self.assertEqual(result['items'][0]['status'],'invalid_submission')
        self.assertTrue(result['items'][0]['changed_since_previous'])

    def test_unchanged_evidence_has_no_incremental_targets(self):
        self.submit();prior=self.root/'prior.json';workflow.write(prior,self.check())
        self.assertEqual(self.check(prior)['changed_items'],[])

    def test_single_opinion_produces_single_incremental_target(self):
        prior=self.root/'prior.json';workflow.write(prior,self.check());self.submit();result=self.check(prior)
        self.assertEqual(result['changed_items'],[self.package['items'][0]['item_id']])

    def test_evidence_path_escape_is_rejected(self):
        self.submit();rows=self.rows();rows[0]['evidence_paths']='["../external.txt"]';self.save_rows(rows)
        self.assertEqual(self.check()['items'][0]['status'],'invalid_submission')

    def test_empty_workbook_caches_do_not_become_valid_evidence(self):
        self.submit(evidence='blank.xlsx');workbook=Workbook();workbook.active['A1']='=1+1';workbook.save(self.root/'evidence/submissions/blank.xlsx')
        rows=self.rows();rows[0]['evidence_sha256']=json.dumps([workflow.file_hash(self.root/'evidence/submissions/blank.xlsx')]);self.save_rows(rows)
        item=self.check()['items'][0];self.assertEqual(item['status'],'invalid_submission')
        self.assertEqual(item['checks'][0]['missing_caches'],['Sheet!A1'])

    def test_incomplete_inventory_rejected(self):
        self.review['fields']=self.review['fields'][:1];workflow.write(self.root/'config/report-field-review.json',self.review)
        workflow.export_package(self.root,self.destination)
        with self.assertRaisesRegex(ValueError,'every current template variable'):self.check()

    def test_preview_fills_only_markers_and_never_scores(self):
        before=(self.root/'templates/report.docx').read_bytes()
        result=preview.create_review_preview(self.spec,{'report_name':'A<&','score':100},self.root/'output',lambda p:self.root/p)
        document=Document(result['report']);text='\n'.join(p.text for p in document.paragraphs)
        self.assertIn('A<&',text);self.assertNotIn('用户补录，未核验',text)
        self.assertIn('Score: 未提供',text);self.assertNotIn('100',text)
        self.assertNotIn('工程核对稿',text)
        audit=workflow.read(result['audit']);self.assertFalse(audit['scoring_performed']);self.assertEqual(audit['missing_fields'],['score'])
        self.assertEqual(before,(self.root/'templates/report.docx').read_bytes())
        self.assertEqual({p.name for p in Path(result['bundle']).iterdir()},{'engineering-review.docx','engineering-review-audit.json'})

    def test_preview_failure_rolls_back_pair(self):
        with patch.object(preview,'render_review_template',side_effect=RuntimeError('render failed')):
            with self.assertRaisesRegex(RuntimeError,'render failed'):
                preview.create_review_preview(self.spec,{},self.root/'output',lambda p:self.root/p)
        self.assertEqual(list((self.root/'output').iterdir()),[])

    def test_preview_requires_template_and_inventory_identity(self):
        self.review['template_sha256']='0'*64;workflow.write(self.root/'config/report-field-review.json',self.review)
        with self.assertRaisesRegex(ValueError,'清单与模板'):preview.create_review_preview(self.spec,{},self.root/'output',lambda p:self.root/p)


if __name__=='__main__':unittest.main()
