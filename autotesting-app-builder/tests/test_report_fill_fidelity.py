"""Report rendering may replace mapped text only."""
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
import sys
import unittest

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "assets/windows-app-template/src"))
from review_preview import render_report_template


class ReportFillFidelityTests(unittest.TestCase):
    def test_only_marker_text_changes(self):
        with TemporaryDirectory() as folder:
            source = Path(folder) / "template.docx"
            result = Path(folder) / "report.docx"
            doc = Document()
            doc.sections[0].header.paragraphs[0].text = "固定页眉"
            doc.sections[0].footer.paragraphs[0].text = "固定页脚"
            doc.add_paragraph("固定正文")
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "得分：{{ score }}"
            doc.save(source)
            render_report_template(source, result, {"score": "4.5"})

            with ZipFile(source) as before, ZipFile(result) as after:
                self.assertEqual(before.namelist(), after.namelist())
                changed = [name for name in before.namelist()
                           if before.read(name) != after.read(name)]
                self.assertEqual(changed, ["word/document.xml"])
                old = list(etree.fromstring(before.read("word/document.xml")).iter())
                new = list(etree.fromstring(after.read("word/document.xml")).iter())
                self.assertEqual(len(old), len(new))
                edits = 0
                for left, right in zip(old, new):
                    self.assertEqual((left.tag, left.attrib, left.tail),
                                     (right.tag, right.attrib, right.tail))
                    if left.text != right.text:
                        self.assertIn("{{ score }}", left.text)
                        self.assertEqual(right.text, "得分：4.5")
                        edits += 1
                self.assertEqual(edits, 1)


if __name__ == "__main__":
    unittest.main()
