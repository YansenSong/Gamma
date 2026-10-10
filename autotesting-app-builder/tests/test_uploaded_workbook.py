"""Uploaded-workbook contract; recalculation is a test double, not Excel evidence."""
from __future__ import annotations
import importlib.util
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from openpyxl import Workbook, load_workbook

APP=Path(__file__).resolve().parents[1]/'assets/windows-app-template/src'
sys.path.insert(0,str(APP))
import app_service as service


class UploadedWorkbookTest(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        (self.root/'config').mkdir();(self.root/'templates').mkdir()
        book=Workbook();sheet=book.active;sheet.title='Score';sheet['A2']=0;sheet['B2']='=IF(A2<=10,2,0)'
        book.save(self.root/'templates/score.xlsx');book.close()
        doc=Document();doc.add_paragraph('Score: {{score}}');doc.save(self.root/'templates/report.docx')
        spec={'schema_version':'autotesting-project-v1','application':{'id':'demo','name':'Demo','version':'1'},
              'release_status':'business_validated','validation_state':{'status':'approved'},
              'sources':[{'role':'workbook','sha256':service.digest(self.root/'templates/score.xlsx')},
                         {'role':'report_template','sha256':service.digest(self.root/'templates/report.docx')}],
              'scoring':{'mode':'excel_native','formula_bindings':[{'output':'score','sheet':'Score','cell':'B2','formula':'=IF(A2<=10,2,0)'}]},
              'excel_output':{'template':'templates/score.xlsx','input_cells':{'input':{'sheet':'Score','cell':'A2'}},
                              'output_cells':{'score':{'sheet':'Score','cell':'B2'}}},
              'report_output':{'template':'templates/report.docx','fields':{'score':{'kind':'score','id':'score'}}}}
        (self.root/'config/project.json').write_text(json.dumps(spec),encoding='utf-8')
        self.resource=patch.object(service,'resource_path',side_effect=lambda name:self.root/name)
        self.validation=patch.object(service,'validate_project',return_value={'status':'valid'})
        self.resource.start();self.validation.start();self.addCleanup(self.resource.stop);self.addCleanup(self.validation.stop)

    @staticmethod
    def fake_excel(path,_inputs):
        value=load_workbook(path,data_only=False)['Score']['A2'].value
        score=2 if value<=10 else 0
        replacement=Path(str(path)+'.zip')
        with ZipFile(path) as source, ZipFile(replacement,'w') as dest:
            for item in source.infolist():
                data=source.read(item.filename)
                if item.filename=='xl/worksheets/sheet1.xml':
                    data=data.replace(b'<f>IF(A2&lt;=10,2,0)</f><v></v>',f'<f>IF(A2&lt;=10,2,0)</f><v>{score}</v>'.encode())
                dest.writestr(item,data)
        replacement.replace(path)
        return {'engine':'Microsoft Excel','version':'TEST_DOUBLE_ONLY'}

    def test_two_uploads_report_their_own_scores_without_changing_sources(self):
        with patch.object(service,'_excel_native_recalculate',side_effect=self.fake_excel):
            results=[]
            for value in (10,11):
                upload=self.root/f'filled-{value}.xlsx';shutil.copy2(self.root/'templates/score.xlsx',upload)
                book=load_workbook(upload);book['Score']['A2']=value;book.save(upload);book.close()
                before=service.digest(upload)
                result=service.run_job_from_workbook(upload,self.root/'output')
                self.assertEqual(service.digest(upload),before)
                self.assertEqual(load_workbook(result['workbook'],data_only=True)['Score']['B2'].value,result['scores']['score'])
                self.assertIn(f"Score: {result['scores']['score']}",Document(result['report']).paragraphs[0].text)
                results.append(result['scores']['score'])
            self.assertEqual(results,[2,0])

    def test_tampered_score_formula_blocks_before_recalculation(self):
        upload=self.root/'tampered.xlsx';shutil.copy2(self.root/'templates/score.xlsx',upload)
        book=load_workbook(upload);book['Score']['B2']='=99';book.save(upload);book.close()
        with patch.object(service,'_excel_native_recalculate') as recalc:
            with self.assertRaisesRegex(service.InputValidationError,'模板不符'):
                service.run_job_from_workbook(upload,self.root/'output')
            recalc.assert_not_called()

    def test_missing_recalculated_cache_blocks_report(self):
        upload=self.root/'filled.xlsx';shutil.copy2(self.root/'templates/score.xlsx',upload)
        with patch.object(service,'_excel_native_recalculate',return_value={'engine':'Microsoft Excel','version':'TEST_DOUBLE_ONLY'}):
            with self.assertRaisesRegex(service.InputValidationError,'为空或类型不符'):
                service.run_job_from_workbook(upload,self.root/'output')
        self.assertFalse((self.root/'output/demo-run').exists())

    def test_excel_failure_blocks_report(self):
        upload=self.root/'filled.xlsx';shutil.copy2(self.root/'templates/score.xlsx',upload)
        with patch.object(service,'_excel_native_recalculate',side_effect=service.InputValidationError('Excel 不可用')):
            with self.assertRaisesRegex(service.InputValidationError,'Excel 不可用'):
                service.run_job_from_workbook(upload,self.root/'output')
        self.assertFalse((self.root/'output/demo-run').exists())


if __name__=='__main__':unittest.main()
