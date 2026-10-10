"""Stage formulas, then commit only the exact workbook validated by Microsoft Excel.

The companion PowerShell validator writes the evidence JSON. Python never scores.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import uuid
from zipfile import ZipFile

from lxml import etree as ET
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CELL = re.compile(r"^(?:'((?:[^']|'')+)'!|([A-Za-z_][\w.]*)!)?(\$?[A-Z]{1,3}\$?[1-9]\d*)(?::(\$?[A-Z]{1,3}\$?[1-9]\d*))?$")
BLOCKED = ("xl/vbaProject", "xl/activeX/", "xl/externalLinks/", "xl/pivotTables/", "xl/slicers/", "xl/embeddings/", "xl/queryTables/")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def plan_hash(plan):
    return hashlib.sha256(json.dumps(plan, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def _parts(archive):
    wb = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    result = {}
    for sheet in wb.findall(f".//{{{MAIN}}}sheet"):
        target = targets[sheet.get(f"{{{REL}}}id")]
        result[sheet.get("name")] = target.lstrip("/") if target.startswith("/") else "xl/" + target
    return result


def validate_plan(plan, *, base=None):
    path = (Path(base or Path.cwd()) / plan["workbook"]).resolve()
    if not path.is_file() or path.suffix.lower() != ".xlsx" or not os.access(path, os.W_OK):
        raise ValueError("source must be a writable .xlsx file")
    if sha(path) != plan.get("source_sha256"):
        raise ValueError("source SHA-256 mismatch")
    entries = plan.get("formulas")
    cases = plan.get("boundary_cases")
    if not isinstance(entries, list) or not entries or not isinstance(cases, list) or not cases:
        raise ValueError("formulas and native Excel boundary_cases are required")
    with ZipFile(path) as archive:
        names = archive.namelist()
        if archive.testzip() is not None or any(n.startswith(BLOCKED) or n == "xl/connections.xml" for n in names):
            raise ValueError("unsupported or damaged XLSX archive")
        parts = _parts(archive)
        if any(p not in names for p in parts.values()):
            raise ValueError("missing worksheet XML")
    book = load_workbook(path, data_only=False)
    refs = {}
    targets = set()
    try:
        def dependencies(formula, sheet):
            found = set()
            tokens = Tokenizer(formula).items
            if any(t.type == "ERROR" for t in tokens):
                raise ValueError("formula contains an Excel error")
            for token in tokens:
                if token.type != "OPERAND" or token.subtype != "RANGE":
                    continue
                match = CELL.fullmatch(token.value)
                if not match:
                    raise ValueError(f"unsupported or invalid formula reference: {token.value}")
                dep_sheet = (match.group(1) or match.group(2) or sheet).replace("''", "'")
                if dep_sheet not in book:
                    raise ValueError(f"missing referenced sheet: {dep_sheet}")
                left, right = match.group(3).replace("$", ""), (match.group(4) or match.group(3)).replace("$", "")
                min_col, min_row, max_col, max_row = range_boundaries(f"{left}:{right}")
                if (max_col-min_col+1)*(max_row-min_row+1) > 100000:
                    raise ValueError("formula reference range too large")
                for row in book[dep_sheet].iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
                    found.update((dep_sheet, c.coordinate) for c in row)
            return found
        for entry in entries:
            sheet, address = entry["sheet"], entry["cell"]
            key = (sheet, address)
            if key in targets or sheet not in book or not re.fullmatch(r"[A-Z]{1,3}[1-9]\d*", address):
                raise ValueError(f"invalid or duplicate target: {key}")
            targets.add(key)
            ws = book[sheet]
            cell = ws[address]
            if ws.protection.sheet or cell.value is not None or cell.__class__.__name__ == "MergedCell" or entry.get("original", "missing") is not None:
                raise ValueError(f"protected, merged or occupied target: {key}")
            if not entry.get("inputs") or not entry.get("pdf_refs") or not entry.get("rule_explanation") or not entry.get("boundary_checks"):
                raise ValueError(f"formula provenance incomplete: {key}")
            if not all(isinstance(ref, dict) and isinstance(ref.get("page"), int) and ref["page"] > 0 and ref.get("location") for ref in entry["pdf_refs"]):
                raise ValueError(f"PDF page and location required: {key}")
            formula = entry.get("formula")
            if not isinstance(formula, str) or not formula.startswith("="):
                raise ValueError(f"invalid formula: {key}")
            refs[key] = dependencies(formula, sheet)
        for entry in entries:
            declared = set()
            for reference in entry["inputs"]:
                if not isinstance(reference, str) or "!" not in reference:
                    raise ValueError("input reference must use Sheet!Cell")
                input_sheet, input_cell = reference.rsplit("!", 1)
                if input_sheet not in book or not re.fullmatch(r"[A-Z]{1,3}[1-9]\d*", input_cell):
                    raise ValueError(f"invalid declared input: {reference}")
                declared.add((input_sheet, input_cell))
            if not declared <= refs[(entry["sheet"], entry["cell"])]:
                raise ValueError("declared input is not referenced by formula")
            for dep in refs[(entry["sheet"], entry["cell"])]:
                if dep not in targets and book[dep[0]][dep[1]].value is None and dep not in declared:
                    raise ValueError(f"missing formula input declaration: {dep}")
        pending = {dep for edges in refs.values() for dep in edges if dep not in refs and book[dep[0]][dep[1]].data_type == "f"}
        while pending:
            node = pending.pop()
            if node in refs:
                continue
            refs[node] = dependencies(book[node[0]][node[1]].value, node[0])
            pending.update(dep for dep in refs[node] if dep not in refs and book[dep[0]][dep[1]].data_type == "f")
        for case in cases:
            if not isinstance(case.get("inputs"), list) or not case["inputs"] or not isinstance(case.get("expected"), list) or not case["expected"]:
                raise ValueError("every boundary case needs inputs and expected formula results")
            for item in case["inputs"]:
                if item.get("sheet") not in book or not re.fullmatch(r"[A-Z]{1,3}[1-9]\d*", str(item.get("cell", ""))):
                    raise ValueError("invalid boundary input address")
                if (item["sheet"], item["cell"]) in targets or book[item["sheet"]][item["cell"]].data_type == "f":
                    raise ValueError("boundary input overwrites a formula")
            for item in case["expected"]:
                if (item.get("sheet"), item.get("cell")) not in targets or not isinstance(item.get("value"), (int, float, str)):
                    raise ValueError("boundary expectation must target an authored formula")
        def visit(node, active, done):
            if node in active:
                raise ValueError(f"formula cycle: {node}")
            if node in done:
                return
            active.add(node)
            for dep in refs.get(node, ()):
                if dep in refs:
                    visit(dep, active, done)
            active.remove(node); done.add(node)
        done = set()
        for key in refs:
            visit(key, set(), done)
    finally:
        book.close()
    return path, targets


def _patch(source, staged, entries):
    with ZipFile(source) as old, ZipFile(staged, "w") as new:
        parts = _parts(old)
        grouped = {}
        for item in entries:
            grouped.setdefault(parts[item["sheet"]], []).append(item)
        for info in old.infolist():
            payload = old.read(info.filename)
            if info.filename in grouped:
                root = ET.fromstring(payload)
                sheet_data=root.find(f"{{{MAIN}}}sheetData")
                for item in grouped[info.filename]:
                    cells = root.xpath(".//m:sheetData/m:row/m:c[@r=$address]", namespaces={"m": MAIN}, address=item["cell"])
                    if len(cells)>1:raise ValueError(f"target has duplicate XML cell records: {item['sheet']}!{item['cell']}")
                    if cells:
                        cell=cells[0]
                    else:
                        address=item['cell'];letters,row_number=coordinate_from_string(address)
                        rows=sheet_data.xpath("./m:row[@r=$number]",namespaces={"m":MAIN},number=str(row_number))
                        if rows:row=rows[0]
                        else:
                            row=ET.Element(f"{{{MAIN}}}row",r=str(row_number))
                            following=next((child for child in sheet_data if int(child.get('r','0'))>row_number),None)
                            sheet_data.insert(sheet_data.index(following),row) if following is not None else sheet_data.append(row)
                        cell=ET.Element(f"{{{MAIN}}}c",r=address)
                        column=column_index_from_string(letters)
                        following=next((child for child in row if child.tag==f"{{{MAIN}}}c" and column_index_from_string(coordinate_from_string(child.get('r'))[0])>column),None)
                        row.insert(row.index(following),cell) if following is not None else row.append(cell)
                    if len(cell) or cell.get("t") not in (None, "n"):
                        raise ValueError("target XML cell is not empty")
                    ET.SubElement(cell, f"{{{MAIN}}}f").text = item["formula"][1:]
                payload = ET.tostring(root, xml_declaration=True, encoding="UTF-8")
            new.writestr(info, payload)


def _verify_staged(source, staged, entries):
    with ZipFile(source) as old, ZipFile(staged) as new:
        if new.testzip() is not None or old.namelist() != new.namelist():
            raise ValueError("staged XLSX archive failed integrity check")
        changed = {name for name in old.namelist() if old.read(name) != new.read(name)}
        if changed != {_parts(old)[e["sheet"]] for e in entries}:
            raise ValueError("unexpected XLSX package member changed")
        grouped = {}
        for entry in entries:
            grouped.setdefault(_parts(old)[entry["sheet"]], []).append(entry["cell"])
        for part, addresses in grouped.items():
            prior = ET.fromstring(old.read(part))
            updated = ET.fromstring(new.read(part))
            for address in addresses:
                cells = updated.xpath(".//m:sheetData/m:row/m:c[@r=$address]", namespaces={"m": MAIN}, address=address)
                if len(cells) != 1 or len(cells[0]) != 1 or cells[0][0].tag != f"{{{MAIN}}}f":
                    raise ValueError("unexpected target cell XML changes")
                existing=prior.xpath(".//m:sheetData/m:row/m:c[@r=$address]",namespaces={"m":MAIN},address=address)
                cell=cells[0];cell.remove(cell[0])
                if not existing:
                    row=cell.getparent();row.remove(cell)
                    row_number=coordinate_from_string(address)[1]
                    if len(row)==0 and not prior.xpath(".//m:sheetData/m:row[@r=$number]",namespaces={"m":MAIN},number=str(row_number)):
                        row.getparent().remove(row)
            if ET.tostring(prior, method="c14n") != ET.tostring(updated, method="c14n"):
                raise ValueError("non-target worksheet XML changed")
    book = load_workbook(staged, data_only=False, read_only=True)
    try:
        if any(book[e["sheet"]][e["cell"]].value != e["formula"] for e in entries):
            raise ValueError("staged formula mismatch")
    finally:
        book.close()


def prepare(plan_path, staged_path, manifest_path):
    plan = json.loads(Path(plan_path).read_text(encoding="utf-8"))
    path, _ = validate_plan(plan, base=Path(plan_path).resolve().parent)
    staged = Path(staged_path).resolve()
    if staged == path or staged.exists():
        raise ValueError("stage must be a new path distinct from source")
    staged.parent.mkdir(parents=True, exist_ok=True)
    try:
        _patch(path, staged, plan["formulas"])
        _verify_staged(path, staged, plan["formulas"])
    except Exception:
        staged.unlink(missing_ok=True)
        raise
    manifest = {"status": "staged", "plan_sha256": plan_hash(plan), "source_sha256": sha(path), "staged_sha256": sha(staged), "staged_path": str(staged)}
    Path(manifest_path).write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def commit(plan_path, manifest_path, evidence_path, receipt_path):
    plan = json.loads(Path(plan_path).read_text(encoding="utf-8"))
    source, _ = validate_plan(plan, base=Path(plan_path).resolve().parent)
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    staged = Path(manifest["staged_path"]).resolve()
    if staged == source or not staged.is_file():
        raise ValueError("missing independent staged workbook")
    _verify_staged(source, staged, plan["formulas"])
    if manifest.get("plan_sha256") != plan_hash(plan) or manifest.get("source_sha256") != sha(source) or manifest.get("staged_sha256") != sha(staged):
        raise ValueError("stage manifest no longer matches plan or files")
    evidence_path = Path(evidence_path).resolve()
    receipt_path = Path(receipt_path).resolve()
    if receipt_path == source or receipt_path == staged or evidence_path in (source, staged):
        raise ValueError("evidence and receipt paths must not replace workbooks")
    if evidence_path.exists():
        raise ValueError("evidence output must be a new file")
    validator = Path(__file__).with_name("validate_formula_stage.ps1")
    result = subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(validator),
                             str(Path(plan_path).resolve()), str(Path(manifest_path).resolve()), str(evidence_path)],
                            capture_output=True, text=True, timeout=300, check=False)
    if result.returncode != 0 or not evidence_path.is_file():
        raise ValueError("Microsoft Excel native validation failed: " + (result.stderr or result.stdout).strip()[:500])
    evidence = json.loads(Path(evidence_path).read_text(encoding="utf-8-sig"))
    if (evidence.get("status") != "passed" or evidence.get("engine") != "Microsoft Excel" or not evidence.get("excel_version")
            or evidence.get("plan_sha256") != plan_hash(plan) or evidence.get("source_sha256") != sha(source)
            or evidence.get("staged_sha256") != sha(staged) or len(evidence.get("cases", [])) != len(plan["boundary_cases"])
            or any(case.get("status") != "passed" for case in evidence["cases"])):
        raise ValueError("native Excel evidence does not bind this plan, stage and all boundary cases")
    if (plan_hash(json.loads(Path(plan_path).read_text(encoding="utf-8"))) != manifest["plan_sha256"]
            or sha(source) != manifest["source_sha256"] or sha(staged) != manifest["staged_sha256"]):
        raise ValueError("plan or workbook changed during native Excel validation")
    before = sha(source)
    backup = source.with_name(source.name + f".backup-{uuid.uuid4().hex}")
    with source.open("rb") as inp, backup.open("xb") as out:
        shutil.copyfileobj(inp, out); out.flush(); os.fsync(out.fileno())
    if sha(backup) != before:
        raise ValueError("backup hash mismatch")
    receipt = {"status": "engineering_sample", "original_path": str(source), "original_sha256": before,
               "backup_path": str(backup), "backup_sha256": sha(backup), "new_sha256": sha(staged),
               "plan_sha256": plan_hash(plan), "native_excel_evidence_sha256": sha(evidence_path),
               "modified_cells": [{k: e[k] for k in ("sheet", "cell", "formula", "pdf_refs")} for e in plan["formulas"]]}
    receipt_path.parent.mkdir(parents=True, exist_ok=True)
    if receipt_path.exists():
        raise ValueError("receipt output must be a new file")
    receipt_tmp = receipt_path.with_name(receipt_path.name + f".{uuid.uuid4().hex}.tmp")
    try:
        receipt_tmp.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        if sha(source) != before or sha(staged) != receipt["new_sha256"]:
            raise ValueError("source or staged workbook changed before replacement")
        os.replace(staged, source)
        try:
            check = load_workbook(source, read_only=True, data_only=False)
            try:
                if any(check[e["sheet"]][e["cell"]].value != e["formula"] for e in plan["formulas"]):
                    raise ValueError("post-replacement formula mismatch")
            finally:
                check.close()
            if sha(source) != receipt["new_sha256"]:
                raise ValueError("post-replacement hash mismatch")
            os.replace(receipt_tmp, receipt_path)
        except Exception:
            restore = source.with_name(source.name + f".restore-{uuid.uuid4().hex}")
            shutil.copy2(backup, restore)
            os.replace(restore, source)
            raise
    finally:
        receipt_tmp.unlink(missing_ok=True)
    return receipt


def main():
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    stage = sub.add_parser("prepare"); stage.add_argument("plan", type=Path); stage.add_argument("--staged", required=True, type=Path); stage.add_argument("--manifest", required=True, type=Path)
    write = sub.add_parser("commit"); write.add_argument("plan", type=Path); write.add_argument("--manifest", required=True, type=Path)
    write.add_argument("--evidence", required=True, type=Path); write.add_argument("--receipt", required=True, type=Path)
    args = parser.parse_args()
    result = prepare(args.plan, args.staged, args.manifest) if args.command == "prepare" else commit(args.plan, args.manifest, args.evidence, args.receipt)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
