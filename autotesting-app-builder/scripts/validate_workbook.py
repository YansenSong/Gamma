from __future__ import annotations
import argparse, hashlib, json, sys
from pathlib import Path
from openpyxl import load_workbook
from _bootstrap import REPO_ROOT
sys.path.insert(0,str(REPO_ROOT))
from ivista_engine.authoring.excel_analyzer import _archive_features

def main() -> int:
    parser=argparse.ArgumentParser(description="Check workbook identity, unsupported features, and formula caches.")
    parser.add_argument("workbook",type=Path); parser.add_argument("--expected-sha256"); parser.add_argument("--require-recalculated",action="store_true")
    args=parser.parse_args(); path=args.workbook.resolve()
    if path.suffix.lower()!=".xlsx": print(json.dumps({"status":"error","error":"only .xlsx is supported"})); return 2
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    if args.expected_sha256 and digest!=args.expected_sha256:
        print(json.dumps({"status":"error","error":"source hash mismatch","actualSha256":digest})); return 2
    features=_archive_features(path)
    formulas=load_workbook(path,data_only=False,read_only=True)
    cached=load_workbook(path,data_only=True,read_only=True)
    formula_cells=[]; missing=[]
    for ws in formulas.worksheets:
        cached_ws=cached[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type=='f':
                    item={"sheet":ws.title,"cell":cell.coordinate}
                    formula_cells.append(item)
                    if cached_ws[cell.coordinate].value is None: missing.append(item)
    formulas.close(); cached.close()
    status="passed"
    errors=[]
    if features: errors.append("unsupported_features:"+",".join(features))
    if args.require_recalculated and missing: errors.append("formula_cache_missing")
    if errors: status="failed"
    print(json.dumps({"status":status,"sourceSha256":digest,"formulaCount":len(formula_cells),
                      "missingCachedResults":missing,"unsupportedFeatures":features,"errors":errors,
                      "note":"openpyxl does not calculate formulas; native Excel recalculation must produce and save cached results."},ensure_ascii=False))
    return 0 if status=="passed" else 1
if __name__=="__main__": raise SystemExit(main())
