"""Describe empty formulas without granting a not-applicable exception."""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook


def inspect(source, workbook, policy):
    digest=lambda path:hashlib.sha256(Path(path).read_bytes()).hexdigest()
    if digest(source)!=policy['workbook_sha256']:raise ValueError('Policy source workbook hash mismatch')
    original=load_workbook(source,data_only=False,read_only=True)
    formulas=load_workbook(workbook,data_only=False,read_only=True)
    cached=load_workbook(workbook,data_only=True,read_only=True)
    rows=[]
    try:
        for entry in policy['cells']:
            sheet,cell=entry['sheet'],entry['cell']
            before=original[sheet][cell];after=formulas[sheet][cell];result=cached[sheet][cell]
            selected=cached[sheet][entry['selection_cell']].value
            if before.data_type!='f' or after.data_type!='f' or before.value!=after.value:
                status='formula_missing_or_changed'
            elif result.data_type=='e':status='calculation_error'
            elif result.value is not None:status='result_present'
            elif selected is not None and str(selected).strip():status='missing_result_with_input'
            else:status='missing_result_and_selection_unresolved'
            rows.append({**entry,'formula':after.value,'cached_result':result.value,'selection_value':selected,
                         'inspection_status':status,'not_applicable_authorized':False,
                         'remaining_review':'Exact charter permission, actual selection record, score denominator treatment and content-bound named approval remain necessary.'})
    finally:original.close();formulas.close();cached.close()
    return {'schema_version':'autotesting-empty-formula-inspection-v1','source_sha256':digest(source),
            'workbook_sha256':digest(workbook),'policy_content_sha256':hashlib.sha256(json.dumps(policy,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest(),
            'cells':rows,'strict_cache_validation_unchanged':True,'business_status':'blocked',
            'note':'Empty selection is an observation, not proof that a scoring project is unselected or may be omitted.'}


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('workbook',type=Path)
    parser.add_argument('--source',type=Path,required=True);parser.add_argument('--policy',type=Path,required=True)
    parser.add_argument('--out',type=Path,required=True);args=parser.parse_args()
    try:
        result=inspect(args.source,args.workbook,json.loads(args.policy.read_text(encoding='utf-8')))
        args.out.parent.mkdir(parents=True,exist_ok=True)
        args.out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
        print(json.dumps({'status':'inspected','not_applicable_authorized':0,'cells':len(result['cells'])},ensure_ascii=False))
    except (ValueError,OSError,KeyError) as exc:print(json.dumps({'status':'failed','error':str(exc)},ensure_ascii=False));raise SystemExit(1)
