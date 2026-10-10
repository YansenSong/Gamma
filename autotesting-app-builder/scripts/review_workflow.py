"""Content-bound review packages and read-only, incremental evidence checks.

Receipts record submitted opinions; they never mutate project approvals or facts.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path
from export_field_review import validate_review
from validate_project_spec import validate as validate_spec, approval_content_hash


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(',', ':'), allow_nan=False).encode()).hexdigest()


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write(path, value):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n',
                          encoding='utf-8', newline='\n')


def inside(root, name):
    relative = Path(name)
    if relative.is_absolute() or '..' in relative.parts:
        raise ValueError('Path must stay inside the project: '+str(name))
    result = (root/relative).resolve()
    result.relative_to(root.resolve())
    return result


def build_package(root):
    root = Path(root).resolve()
    spec = read(root/'config/project.json')
    validate_spec(root/'config/project.json')
    items = []
    def add(kind, identifier, subject, question, locations, payload):
        items.append({'item_id':kind+':'+identifier, 'kind':kind, 'subject':subject,
                      'question':question, 'locations':locations, 'payload':payload,
                      'content_sha256':digest(payload), 'blocks_official':True})
    review_path = spec.get('report_review', {}).get('path')
    if review_path:
        review = read(inside(root, review_path)); validate_review(review)
        template = next(x['sha256'] for x in spec['sources'] if x['role']=='report_template')
        if review['template_sha256'] != template:
            raise ValueError('Report review belongs to a different template')
        for field in review['fields']:
            if field['category'] != 'engineering_literal':
                add('field', field['field_id'], field['meaning'], field['question'],
                    field['word_locations'], field)
        for index, question in enumerate(review.get('static_review_questions', []), 1):
            add('template', str(index), question, '请记录审核依据、模板副本变更与负责人。', [], {'question':question})
    coverage = root/'analysis/charter-coverage.json'
    if coverage.exists():
        for row in read(coverage)['rows']:
            add('rule', row['rule_id'], row.get('reason',row['rule_id']),
                '请核验章程口径、工况和原始记录；一致仅表示当前对照结果。',
                row.get('pdf_refs', []), row)
    questions_path = root/'analysis/Q1-Q10.json'
    questions = {x['id']:x for x in read(questions_path)['questions']} if questions_path.exists() else {}
    for decision in spec.get('business_decisions', []):
        payload = dict(questions.get(decision['id'], decision))
        if 'candidate_interpretations' in payload:
            payload['candidate_interpretations']=[x for x in payload['candidate_interpretations'] if x.get('id')!='python_score_authority']
            payload['product_constraint']='Excel formula results are the sole scoring authority; Python scoring has been removed by user instruction.'
        add('decision', decision['id'], payload.get('title',decision.get('reason','')),
            payload.get('required_confirmation','请具名裁决，附章程条文及原始记录；填写意见不改变软件规则。'),
            [payload['charter_evidence']] if payload.get('charter_evidence') else [], payload)
    for index, reason in enumerate(spec.get('validation_state', {}).get('reasons', []), 1):
        add('blocker', str(index), reason, '请提供解决依据；须更新实现并复验后才可申请解除。', [], {'reason':reason})
    policy_path = root/'config/formula-empty-review.json'
    if policy_path.exists():
        for row in read(policy_path)['cells']:
            add('empty_formula', row['sheet']+'!'+row['cell'], row['question'],
                '请确认是否允许不参与：须附章程条文、项目选择记录、分母/总分处理及具名意见。',
                [row['sheet']+'!'+row['cell']], row)
    ids = [x['item_id'] for x in items]
    if len(ids) != len(set(ids)):
        raise ValueError('Duplicate review item ID')
    package = {'schema_version':'autotesting-review-package-v1', 'project_id':spec['application']['id'],
               'project_spec_sha256':file_hash(root/'config/project.json'), 'sources':spec['sources'],
               'approval_content_sha256':{k:approval_content_hash(spec,k) for k in
                 ('scoring_rules','excel_mapping','report_mapping','calculation_authority')},
               'business_status':'blocked', 'counts':dict(Counter(x['kind'] for x in items)), 'items':items}
    package['package_sha256'] = digest(package)
    return package


def export_package(root, destination):
    package = build_package(root)
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    write(destination/'package.json', package)
    keys = ['package_sha256','item_id','content_sha256','subject','question','locations',
            'reviewer','reviewed_at','decision','reason','evidence_paths','evidence_sha256','record_ids']
    with (destination/'opinions.csv').open('w',encoding='utf-8-sig',newline='') as stream:
        writer = csv.DictWriter(stream,fieldnames=keys,lineterminator='\n');writer.writeheader()
        for item in package['items']:
            writer.writerow({'package_sha256':package['package_sha256'],
                             **{k:item[k] for k in ('item_id','content_sha256','subject','question')},
                             'locations':'; '.join(item['locations'])})
    text = ['# 业务审核包', '', '状态：待审核；填写、签署和导入均不授予正式发布许可。', '',
            '逐项填写 opinions.csv：reviewer 为真实审核人；reviewed_at 为带时区的 ISO 日期时间；',
            'decision 可填 accept / reject / needs_evidence / not_applicable；reason 记录依据。',
            'evidence_paths、evidence_sha256、record_ids 分别填 JSON 字符串数组，长度一致。',
            '证据路径相对独立项目根目录，保存于 evidence/submissions/；不得改动原模板或分析记录。',
            '每条证据记录须填写文件 SHA-256 和原始记录编号；缺失则继续阻断。', '',
            '回收前不要编辑 package.json；材料或映射改变后重新导出审核包，旧意见自动失效。',
            '复验只检查材料身份、结构、规范和意见完整性，不能证明原始记录真实或裁决正确。', '',
            '统计：'+json.dumps(package['counts'],ensure_ascii=False), '',
            '|编号|审核事项|待确认|', '|---|---|---|']
    for item in package['items']:
        text.append('|'+ '|'.join(str(item[k]).replace('|','\\|').replace('\n',' ') for k in
                                  ('item_id','subject','question'))+'|')
    (destination/'README.md').write_text('\n'.join(text)+'\n',encoding='utf-8',newline='\n')
    return package


def revalidate(root, package_path, opinions_path, previous=None):
    from datetime import datetime
    root = Path(root).resolve(); package = read(package_path)
    supplied = dict(package); supplied.pop('package_sha256', None)
    if digest(supplied) != package.get('package_sha256'):
        raise ValueError('Review package content hash mismatch')
    current = build_package(root)
    if current['package_sha256'] != package['package_sha256']:
        raise ValueError('Review package is stale; re-export before collecting new opinions')
    spec=read(root/'config/project.json')
    field_review=read(inside(root,spec['report_review']['path'])) if spec.get('report_review') else {'fields':[]}
    # Explicitly compare the inventory with the current template, including merged-cell XML.
    from export_field_review import docx_variable_locations
    template_fields=set(docx_variable_locations(inside(root,spec['report_output']['template'])))
    inventory_fields={x['field_id'] for x in field_review['fields']}
    if template_fields!=inventory_fields:
        raise ValueError('Report inventory does not cover every current template variable')
    prior = read(previous) if previous else {}
    prior_rows = {x['item_id']:x for x in prior.get('items', [])} if prior.get('package_sha256') == package['package_sha256'] else {}
    with Path(opinions_path).open(encoding='utf-8-sig',newline='') as stream:
        reader=csv.DictReader(stream)
        required={'package_sha256','item_id','content_sha256','reviewer','reviewed_at','decision','reason',
                  'evidence_paths','evidence_sha256','record_ids'}
        if not required<=set(reader.fieldnames or []):raise ValueError('Opinion CSV columns are incomplete')
        opinions=list(reader)
    by_id = {}
    expected = {x['item_id']:x for x in package['items']}
    for row in opinions:
        if None in row or any(value is None for value in row.values()):
            raise ValueError('Malformed opinion CSV row')
        identifier = row.get('item_id')
        if identifier not in expected or identifier in by_id:
            raise ValueError('Unknown or duplicate opinion: '+str(identifier))
        item = expected[identifier]
        if row.get('package_sha256') != package['package_sha256'] or row.get('content_sha256') != item['content_sha256']:
            raise ValueError('Opinion content binding mismatch: '+identifier)
        by_id[identifier] = row
    results = []
    for item in package['items']:
        row = by_id.get(item['item_id'], {})
        errors = []; evidence = []; checks = []
        populated = any(str(row.get(k,'')).strip() for k in
                        ('reviewer','reviewed_at','decision','reason','evidence_paths','evidence_sha256','record_ids'))
        if populated:
            for key in ('reviewer','reviewed_at','reason'):
                if not row.get(key,'').strip():errors.append('missing_'+key)
            if row.get('decision') not in ('accept','reject','needs_evidence','not_applicable'):
                errors.append('invalid_decision')
            try:
                moment = datetime.fromisoformat(row.get('reviewed_at',''))
                if moment.tzinfo is None:errors.append('review_time_requires_timezone')
            except ValueError:errors.append('invalid_review_time')
            try:
                paths, hashes, records = [json.loads(row.get(k) or '[]') for k in
                                         ('evidence_paths','evidence_sha256','record_ids')]
                if not all(isinstance(x,list) for x in (paths,hashes,records)) or not (len(paths)==len(hashes)==len(records)):
                    raise ValueError('Evidence arrays must have equal length')
                if not paths:errors.append('missing_evidence')
                for name, sha, record in zip(paths,hashes,records):
                    relative=Path(name)
                    if relative.parts[:2] != ('evidence','submissions'):
                        raise ValueError('New evidence must be under evidence/submissions/')
                    path = inside(root,name)
                    if not isinstance(record,str) or not record.strip():errors.append('missing_record_id')
                    if not path.is_file():errors.append('evidence_missing:'+name);continue
                    actual = file_hash(path)
                    evidence.append({'path':name,'expected_sha256':sha,'actual_sha256':actual,'record_id':record})
                    if actual!=sha:errors.append('evidence_hash_mismatch:'+name)
            except (ValueError,TypeError,OSError) as exc:errors.append('invalid_evidence:'+str(exc))
        # Reinspect supplied evidence; prior records only identify changes, never bypass checks.
        if populated and not errors:
            for entry in evidence:
                path=inside(root,entry['path'])
                if path.suffix.lower()=='.xlsx':
                    from openpyxl import load_workbook
                    try:
                        formulas=load_workbook(path,data_only=False,read_only=True)
                        cached=load_workbook(path,data_only=True,read_only=True)
                        blanks=[];excel_errors=[]
                        try:
                            for sheet in formulas:
                                for cells in sheet.iter_rows():
                                    for cell in cells:
                                        value=cached[sheet.title][cell.coordinate]
                                        if cell.data_type=='f' and value.value is None:blanks.append(sheet.title+'!'+cell.coordinate)
                                        if value.data_type=='e':excel_errors.append(sheet.title+'!'+cell.coordinate)
                        finally:formulas.close();cached.close()
                        checks.append({'path':entry['path'],'type':'workbook_cache','status':'failed' if blanks or excel_errors else 'structure_checked',
                                       'missing_caches':blanks,'excel_errors':excel_errors,
                                       'native_recalculation_verified':False})
                    except Exception as exc:errors.append('workbook_unreadable:'+str(exc))
                elif path.suffix.lower()=='.docx':
                    from docx import Document
                    try:Document(path);checks.append({'path':entry['path'],'type':'docx_structure','status':'structure_checked'})
                    except Exception as exc:errors.append('report_unreadable:'+str(exc))
                else:checks.append({'path':entry['path'],'type':'file_identity','status':'hash_checked','semantic_validation':'pending'})
            if any(x['status']=='failed' for x in checks):errors.append('evidence_structure_failed')
        fingerprint=digest({'item':item,'opinion':row,'evidence':evidence,'errors':errors,'checks':checks})
        changed=prior_rows.get(item['item_id'],{}).get('fingerprint') != fingerprint
        followups={'field':['source_semantics_units_and_trial_review','approved_report_mapping_revalidation'],
                   'rule':['trial_fact_contract_validation','project_specific_charter_comparison'],
                   'decision':['named_business_decision','decision_to_excel_formula_validation'],
                   'blocker':['project_specific_business_review'],
                   'template':['approved_copy_layout_and_mapping_review'],
                   'empty_formula':['charter_permission_and_selection_record_review','strict_formula_cache_validation']}[item['kind']]
        results.append({'item_id':item['item_id'],'kind':item['kind'],
                        'status':'pending' if not populated else 'invalid_submission' if errors else 'opinion_submitted',
                        'fingerprint':fingerprint,'changed_since_previous':changed,
                        'opinion':row,'evidence':evidence,'checks':checks,'errors':errors,
                        'followup_checks_pending':followups,'blocks_official':True})
    return {'schema_version':'autotesting-evidence-revalidation-v1','package_sha256':package['package_sha256'],
            'project_spec_validation':validate_spec(root/'config/project.json'),
            'opinion_file_sha256':file_hash(opinions_path), 'business_status':'blocked',
            'report_inventory_validation':{'status':'complete','template_variables':len(template_fields)},
            'official_output_allowed':False,'counts':dict(Counter(x['status'] for x in results)),
            'changed_items':[x['item_id'] for x in results if x['changed_since_previous']],
            'items':results, 'note':'Submitted opinions are not approvals. Evidence identity/structure does not establish trial truth. No project files were changed.'}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    commands=parser.add_subparsers(dest='command',required=True)
    export=commands.add_parser('export');export.add_argument('project_root',type=Path);export.add_argument('--out',type=Path,required=True)
    check=commands.add_parser('revalidate');check.add_argument('project_root',type=Path)
    check.add_argument('--package',type=Path,required=True);check.add_argument('--opinions',type=Path,required=True)
    check.add_argument('--previous',type=Path);check.add_argument('--out',type=Path,required=True)
    args=parser.parse_args()
    try:
        if args.command=='export':result=export_package(args.project_root,args.out)
        else:
            result=revalidate(args.project_root,args.package,args.opinions,args.previous);write(args.out,result)
        print(json.dumps({'status':'written','business_status':'blocked','counts':result['counts']},ensure_ascii=False))
        return 1 if result.get('counts',{}).get('invalid_submission') else 0
    except (ValueError,KeyError,OSError) as exc:
        print(json.dumps({'status':'failed','error':str(exc),'official_output_allowed':False},ensure_ascii=False));return 1


if __name__=='__main__':raise SystemExit(main())
