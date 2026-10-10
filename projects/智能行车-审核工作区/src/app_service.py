"""Standalone scoring, protected workbook, DOCX, and atomic audit bundle service."""
from __future__ import annotations
import hashlib,json,shutil,subprocess,sys,tempfile,uuid,math
from datetime import datetime,timezone
from pathlib import Path
from docxtpl import DocxTemplate
from openpyxl import load_workbook
from input_validation import validate_facts,InputValidationError
from project_validation import validate as validate_project,docx_markers
from report_safety import mark_engineering_copy, requires_draft

def resource_path(name): return Path(getattr(sys,"_MEIPASS",Path(__file__).resolve().parents[1]))/name
def load_spec(): return json.loads(resource_path("config/project.json").read_text(encoding="utf-8"))
def digest(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def plain(value): return float(value) if hasattr(value,"as_tuple") else value
def cell_signature(cell):
    value=cell.value; blank=value is None or value==""
    return (None if blank else value,"blank" if blank else cell.data_type,cell.number_format,cell.protection.locked,
            cell.comment.text if cell.comment else None,cell.hyperlink.target if cell.hyperlink else None)
def workbook_snapshot(book,targets):
    snapshot={"sheets":tuple(book.sheetnames),"defined_names":tuple(sorted((name,entry.attr_text) for name,entry in book.defined_names.items())),"cells":{}}
    for sheet in book.worksheets:
        snapshot[f"merged:{sheet.title}"]=tuple(sorted(str(r) for r in sheet.merged_cells.ranges))
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None or cell.has_style or cell.comment or cell.hyperlink:
                    key=(sheet.title,cell.coordinate)
                    if key not in targets: snapshot["cells"][key]=cell_signature(cell)
    return snapshot
def mapping_targets(book,mappings,*,formula_outputs=False):
    targets=set()
    for key,mapping in mappings.items():
        sheet=mapping.get("sheet"); address=mapping.get("cell")
        if sheet not in book.sheetnames or not isinstance(address,str): raise InputValidationError(f"Excel 映射无效：{key}")
        cell=book[sheet][address]
        if cell.coordinate!=address.upper(): raise InputValidationError(f"Excel 单元格地址无效：{key}")
        marker=(sheet,cell.coordinate)
        if marker in targets: raise InputValidationError(f"多个 Excel 字段映射到同一单元格：{sheet}!{cell.coordinate}")
        if any(cell.coordinate in merged and cell.coordinate!=merged.start_cell.coordinate for merged in book[sheet].merged_cells.ranges):
            raise InputValidationError(f"Excel 映射指向合并区域非左上角：{sheet}!{cell.coordinate}")
        if not formula_outputs and book[sheet].protection.sheet and cell.protection.locked: raise InputValidationError(f"Excel 目标单元格受保护：{sheet}!{cell.coordinate}")
        if formula_outputs and cell.data_type!='f':
            raise InputValidationError(f"Excel 原生评分结果单元格必须包含公式：{sheet}!{cell.coordinate}")
        if not formula_outputs and isinstance(cell.value,str) and cell.value.startswith("="):
            raise InputValidationError(f"拒绝覆盖 Excel 公式：{sheet}!{cell.coordinate}")
        targets.add(marker)
    return targets
def report_context(fields,facts,scores,system=None):
    context={}
    for name,source in fields.items():
        kind,key=source.get("kind"),source.get("id")
        if kind=="input" and source.get('optional'): context[name]=plain(facts.get(key)) if facts.get(key) not in (None,'') else '未提供'
        elif kind=="input" and key in facts: context[name]=plain(facts[key])
        elif kind=="score" and key in scores: context[name]=plain(scores[key])
        elif kind=="system" and key in (system or {}): context[name]=(system or {})[key]
        elif kind=="literal": context[name]=source.get("value","")
        else: raise InputValidationError(f"报告字段 {name} 的数据来源未定义")
    return context
def load_facts_from_excel(path,spec=None):
    spec=spec or load_spec(); cfg=spec.get("excel_input",{}); mappings=cfg.get("field_cells",{})
    book=load_workbook(path,data_only=True,read_only=True)
    try:
        facts={field["id"]:field["default"] for field in spec["scoring"]["inputs"] if "default" in field}
        facts.update({key:book[item["sheet"]][item["cell"]].value for key,item in mappings.items()})
    except (KeyError,ValueError) as exc: raise InputValidationError(f"无法从 Excel 导入测评字段：{exc}") from exc
    finally: book.close()
    return validate_facts(spec,facts)
def _excel_native_recalculate(workbook_path,input_values):
    script=resource_path("scripts/recalculate_excel.ps1")
    payload=Path(tempfile.gettempdir())/f"autotesting-excel-input-{uuid.uuid4().hex}.json"
    payload.write_text(json.dumps(input_values,ensure_ascii=False),encoding="utf-8")
    try:
        result=subprocess.run(["powershell.exe","-NoProfile","-NonInteractive","-ExecutionPolicy","Bypass","-File",str(script),str(workbook_path),str(payload)],
            capture_output=True,text=True,timeout=120,check=False,creationflags=getattr(subprocess,"CREATE_NO_WINDOW",0))
    except (OSError,subprocess.TimeoutExpired) as exc: raise InputValidationError(f"Excel 原生重算失败：{exc}") from exc
    finally: payload.unlink(missing_ok=True)
    if result.returncode: raise InputValidationError("Excel 原生重算失败："+(result.stderr.strip() or result.stdout.strip()))
    try: return json.loads(result.stdout.strip().splitlines()[-1])
    except (IndexError,json.JSONDecodeError) as exc: raise InputValidationError("Excel 未返回有效重算记录") from exc
def run_job_from_facts(facts_raw,output_dir):
    spec=load_spec(); mode=spec['scoring'].get('mode')
    if mode!='excel_native':raise InputValidationError('仅支持 Excel 原生公式评分，禁止 Python 补算')
    validate_project(resource_path('config/project.json'))
    facts=validate_facts(spec,facts_raw);scores={}
    app=spec["application"]; xcfg=spec["excel_output"]; workbook_template=resource_path(xcfg["template"])
    expected=next((x["sha256"] for x in spec["sources"] if x["role"]=="workbook"),None)
    if not expected or digest(workbook_template)!=expected: raise InputValidationError("Excel 模板 SHA-256 与项目配置不符")
    root=Path(output_dir).resolve(); root.mkdir(parents=True,exist_ok=True)
    stage=root/f".{app['id']}-{uuid.uuid4().hex}.tmp"; stage.mkdir()
    try:
        xlsx=stage/f"{app['id']}-score.xlsx"; shutil.copy2(workbook_template,xlsx)
        book=load_workbook(xlsx,data_only=False)
        input_targets=mapping_targets(book,xcfg.get("input_cells",{}))
        formula_outputs=True
        if not xcfg.get("output_cells"):raise InputValidationError("Excel 评分输出映射尚未确认")
        output_targets=mapping_targets(book,xcfg.get("output_cells",{}),formula_outputs=True)
        if input_targets & output_targets: raise InputValidationError("Excel 输入与结果映射不能覆盖同一单元格")
        protected_targets=input_targets
        protected=workbook_snapshot(book,protected_targets)
        book.close()
        input_values={f"{item['sheet']}!{item['cell']}":plain(facts[key]) for key,item in xcfg.get("input_cells",{}).items()}
        recalc=_excel_native_recalculate(xlsx,input_values)
        if recalc.get("engine")!="Microsoft Excel" or not recalc.get("version"):raise InputValidationError("缺少 Microsoft Excel 原生重算证据")
        output_book=load_workbook(xlsx,data_only=True,read_only=True)
        try:
            for key,item in xcfg.get('output_cells',{}).items():
                cell=output_book[item['sheet']][item['cell']]
                if cell.data_type=='e':raise InputValidationError('Excel 结果含计算错误：'+key)
                scores[key]=cell.value
        finally: output_book.close()
        for key,value in scores.items():
            expected_type=xcfg.get("output_cells",{}).get(key,{}).get("value_type","number")
            valid=(isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value)) if expected_type=="number" else isinstance(value,str) and bool(value.strip())
            if not valid: raise InputValidationError(f"Excel 原生重算后结果 {key} 为空或类型不符")
        check=load_workbook(xlsx,data_only=False,read_only=False)
        try:
            current=workbook_snapshot(check,protected_targets)
            if current!=protected:
                changed=[f"{key}: {protected.get(key)!r} -> {current.get(key)!r}" for key in (set(protected)|set(current))-{'cells'} if protected.get(key)!=current.get(key)]
                changed.extend(f"cell {key}: {protected['cells'].get(key)!r} -> {current['cells'].get(key)!r}" for key in set(protected['cells'])|set(current['cells']) if protected['cells'].get(key)!=current['cells'].get(key))
                raise InputValidationError("Excel 写入改变了未授权目标之外的模板内容："+"; ".join(changed[:8]))
        finally: check.close()
        rcfg=spec["report_output"]; report_template=resource_path(rcfg["template"])
        report_expected=next((x["sha256"] for x in spec["sources"] if x["role"]=="report_template"),None)
        if not report_expected or digest(report_template)!=report_expected:raise InputValidationError("报告模板 SHA-256 与项目配置不符")
        audit_id=uuid.uuid4().hex
        context=report_context(rcfg["fields"],facts,scores,{"excel_version":recalc.get("version") or "未记录","audit_id":audit_id})
        draft=requires_draft(spec)
        policy=spec.get('report_safety',{'provenance':'工程样例数据，业务规则与原始试验来源未获正式审批。'})
        if draft:
            for key in policy.get('legacy_result_fields',[]):
                context[key]=str(context[key])+policy.get('result_suffix','（工程参考值）')
        markers=docx_markers(report_template)["names"]
        if markers!=set(rcfg["fields"]):
            undeclared=sorted(markers-set(rcfg["fields"]));unmarked=sorted(set(rcfg["fields"])-markers)
            details=[text for text in ("模板中未声明的标记："+", ".join(undeclared) if undeclared else "",
                                       "无对应模板标记的报告字段："+", ".join(unmarked) if unmarked else "") if text]
            raise InputValidationError("报告模板填充标记与配置不一致："+"; ".join(details))
        docx=stage/f"{app['id']}-report.docx"; document=DocxTemplate(str(report_template)); document.render(context); document.save(str(docx))
        visible_status=mark_engineering_copy(docx,policy) if draft else {'status':'business_validated'}
        audit={"schema_version":"autotesting-run-audit-v1","release_status":spec.get("release_status","engineering_sample"),"visible_report_status":visible_status,
          "project_id":app["id"],"audit_id":audit_id,"timestamp_utc":datetime.now(timezone.utc).isoformat(timespec="seconds"),
          "project_spec_sha256":digest(resource_path("config/project.json")),"workbook_template_sha256":digest(workbook_template),
          "report_template_sha256":digest(report_template),"facts":{k:plain(v) for k,v in facts.items()},"scores":{k:plain(v) for k,v in scores.items()},
          "calculation_authority":{"mode":mode,"recalculation":recalc},"output_workbook":{"path":xlsx.name,"sha256":digest(xlsx)},
          "output_report":{"path":docx.name,"sha256":digest(docx)}}
        audit_path=stage/f"{app['id']}-audit.json"; audit_path.write_text(json.dumps(audit,ensure_ascii=False,sort_keys=True,indent=2)+"\n",encoding="utf-8",newline="\n")
        final=root/f"{app['id']}-run"; index=2
        while final.exists(): final=root/f"{app['id']}-run-{index}"; index+=1
        stage.replace(final)
        return {"bundle":final,"workbook":final/xlsx.name,"report":final/docx.name,"audit":final/audit_path.name,"scores":scores}
    except Exception:
        shutil.rmtree(stage,ignore_errors=True); raise
def run_job(input_path,output_dir):
    source=Path(input_path).resolve()
    try: facts=json.loads(source.read_text(encoding="utf-8-sig"))
    except (OSError,json.JSONDecodeError) as exc: raise InputValidationError(f"无法读取 JSON 输入：{exc}") from exc
    return run_job_from_facts(facts,output_dir)

def run_job_from_workbook(input_path, output_dir):
    """Recalculate a copy of the supplied, filled workbook and report its own scores."""
    source=Path(input_path).resolve()
    if source.suffix.lower()!='.xlsx' or not source.is_file():
        raise InputValidationError('请选择已填写的 .xlsx 评分工作簿')
    spec=load_spec()
    if spec.get('scoring',{}).get('mode')!='excel_native':
        raise InputValidationError('仅支持 Excel 原生公式评分')
    validate_project(resource_path('config/project.json'))
    xcfg=spec['excel_output']; outputs=xcfg.get('output_cells',{})
    if not outputs: raise InputValidationError('Excel 评分输出映射尚未确认')
    template=resource_path(xcfg['template'])
    expected=next((item['sha256'] for item in spec['sources'] if item['role']=='workbook'),None)
    if not expected or digest(template)!=expected: raise InputValidationError('Excel 模板 SHA-256 与项目配置不符')
    template_book=load_workbook(template,data_only=False)
    upload=load_workbook(source,data_only=False)
    try:
        if template_book.sheetnames!=upload.sheetnames:
            raise InputValidationError('上传工作簿的工作表结构与模板不符')
        input_cells=xcfg.get('input_cells',{})
        allowed=mapping_targets(template_book,input_cells)
        mapped_outputs=mapping_targets(template_book,outputs,formula_outputs=True)
        if allowed & mapped_outputs: raise InputValidationError('Excel 输入与结果映射不能重叠')
        for sheet in template_book.worksheets:
            other=upload[sheet.title]
            if tuple(sorted(str(x) for x in sheet.merged_cells.ranges))!=tuple(sorted(str(x) for x in other.merged_cells.ranges)):
                raise InputValidationError('上传工作簿的合并单元格与模板不符：'+sheet.title)
            coordinates={cell.coordinate for row in sheet for cell in row if cell.value is not None or cell.has_style or cell.comment or cell.hyperlink}
            coordinates.update(cell.coordinate for row in other for cell in row if cell.value is not None or cell.has_style or cell.comment or cell.hyperlink)
            for address in coordinates:
                key=(sheet.title,address)
                if key in allowed:
                    entered=other[address]
                    if entered.data_type in {'f','e'}:
                        raise InputValidationError(f'输入单元格不能是公式或错误：{sheet.title}!{address}')
                    continue
                if cell_signature(sheet[address])!=cell_signature(other[address]):
                    raise InputValidationError(f'上传工作簿与模板不符：{sheet.title}!{address}')
        if tuple(sorted((name,entry.attr_text) for name,entry in template_book.defined_names.items()))!=tuple(sorted((name,entry.attr_text) for name,entry in upload.defined_names.items())):
            raise InputValidationError('上传工作簿的定义名称与模板不符')
    finally:
        template_book.close();upload.close()
    original_sha=digest(source)
    root=Path(output_dir).resolve();root.mkdir(parents=True,exist_ok=True)
    app=spec['application'];stage=root/f".{app['id']}-{uuid.uuid4().hex}.tmp";stage.mkdir()
    try:
        xlsx=stage/f"{app['id']}-score.xlsx";shutil.copy2(source,xlsx)
        recalc=_excel_native_recalculate(xlsx,{})
        if recalc.get('engine')!='Microsoft Excel' or not recalc.get('version'):
            raise InputValidationError('缺少 Microsoft Excel 原生重算证据')
        if digest(source)!=original_sha: raise InputValidationError('上传原件在处理期间发生变化')
        formula_book=load_workbook(xlsx,data_only=False)
        try:
            uploaded_again=load_workbook(source,data_only=False)
            try:
                if workbook_snapshot(formula_book,allowed)!=workbook_snapshot(uploaded_again,allowed):
                    raise InputValidationError('Excel 重算改变了评分工作簿的公式或非输入内容')
            finally:uploaded_again.close()
            for binding in spec['scoring']['formula_bindings']:
                cell=formula_book[binding['sheet']][binding['cell']]
                if cell.data_type!='f' or cell.value!=binding['formula']:
                    raise InputValidationError('重算后评分公式发生变化：'+binding['output'])
        finally:formula_book.close()
        values=load_workbook(xlsx,data_only=True,read_only=True)
        try:
            scores={}
            for key,item in outputs.items():
                cell=values[item['sheet']][item['cell']]
                value=cell.value
                expected_type=item.get('value_type','number')
                valid=(isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value)) if expected_type=='number' else isinstance(value,str) and bool(value.strip())
                if cell.data_type=='e' or not valid:raise InputValidationError(f'Excel 原生重算后结果 {key} 为空或类型不符')
                scores[key]=value
            facts={key:values[item['sheet']][item['cell']].value for key,item in input_cells.items()}
        finally:values.close()
        rcfg=spec['report_output'];report_template=resource_path(rcfg['template'])
        for name,field in rcfg['fields'].items():
            if field.get('kind')=='input' and facts.get(field.get('id')) in (None,''):
                raise InputValidationError('报告字段缺少上传工作簿来源：'+name)
        report_expected=next((x['sha256'] for x in spec['sources'] if x['role']=='report_template'),None)
        if not report_expected or digest(report_template)!=report_expected:raise InputValidationError('报告模板 SHA-256 与项目配置不符')
        audit_id=uuid.uuid4().hex
        context=report_context(rcfg['fields'],facts,scores,{'excel_version':recalc['version'],'audit_id':audit_id})
        draft=requires_draft(spec)
        policy=spec.get('report_safety',{'provenance':'工程样例数据，业务规则与原始试验来源未获正式审批。'})
        if draft:
            for key in policy.get('legacy_result_fields',[]):
                context[key]=str(context[key])+policy.get('result_suffix','（工程参考值）')
        if docx_markers(report_template)['names']!=set(rcfg['fields']):raise InputValidationError('报告模板标记与配置不一致')
        docx=stage/f"{app['id']}-report.docx";document=DocxTemplate(str(report_template));document.render(context);document.save(str(docx))
        visible_status=mark_engineering_copy(docx,policy) if draft else {'status':'business_validated'}
        audit={'schema_version':'autotesting-run-audit-v1','release_status':spec.get('release_status','engineering_sample'),
               'visible_report_status':visible_status,'project_id':app['id'],'audit_id':audit_id,
               'timestamp_utc':datetime.now(timezone.utc).isoformat(timespec='seconds'),
               'project_spec_sha256':digest(resource_path('config/project.json')),
               'input_workbook':{'path':str(source),'sha256':original_sha},
               'workbook_template_sha256':digest(template),'report_template_sha256':digest(report_template),
               'facts':{key:plain(value) for key,value in facts.items()},'scores':{key:plain(value) for key,value in scores.items()},
               'calculation_authority':{'mode':'excel_native','recalculation':recalc},
               'output_workbook':{'path':xlsx.name,'sha256':digest(xlsx)},
               'output_report':{'path':docx.name,'sha256':digest(docx)}}
        audit_path=stage/f"{app['id']}-audit.json";audit_path.write_text(json.dumps(audit,ensure_ascii=False,sort_keys=True,indent=2)+'\n',encoding='utf-8')
        final=root/f"{app['id']}-run";index=2
        while final.exists():final=root/f"{app['id']}-run-{index}";index+=1
        stage.replace(final)
        return {'bundle':final,'workbook':final/xlsx.name,'report':final/docx.name,'audit':final/audit_path.name,'scores':scores}
    except Exception:
        shutil.rmtree(stage,ignore_errors=True);raise
